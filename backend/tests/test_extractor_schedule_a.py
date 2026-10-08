import asyncio
import unittest

from datetime import datetime
from io import BytesIO
from pathlib import Path
import sys
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.models import (
    DocumentType,
    FieldRule,
    FieldRuleMappingMode,
    FormType,
    NormalizedExtractionField,
    NormalizedExtractionResult,
    ScheduleABrokerMoneyRow,
    ScheduleABrokerRow,
)
from app.services.extractor import (
    ExtractionService,
    _extract_fields_from_pages,
    extract_cigna_schedule_a_broker_rows,
    extract_cigna_schedule_a_fields,
    extract_aig_schedule_a_fields,
    extract_aig_broker_rows,
    prefer_authoritative_aig_fields,
    extract_anthem_schedule_a_fields,
    extract_anthem_broker_rows,
    extract_aflac_schedule_a_fields,
    extract_aflac_broker_rows,
    prefer_authoritative_colonial_fields,
    extract_columnar_broker_compensation_rows,
    extract_bcbs_michigan_addendum_broker_rows,
    extract_bcbs_michigan_schedule_a_fields,
    extract_bcbs_michigan_schedule_a_summaries,
    extract_eyemed_broker_rows,
    extract_eyemed_schedule_a_fields,
    extract_eyemed_schedule_a_summaries,
    extract_vsp_schedule_a_fields,
    extract_vsp_broker_rows,
    prefer_authoritative_cigna_summary_fields,
    extract_explicit_benefit_indicator_fields,
    extract_fields_from_groundx_xray,
    extract_hmsa_schedule_a_fields,
    extract_hartford_broker_rows,
    extract_hartford_schedule_a_fields,
    extract_allone_eap_schedule_a_fields,
    extract_american_heritage_broker_rows,
    extract_american_heritage_schedule_a_fields,
    extract_guardian_broker_rows,
    extract_guardian_schedule_a_fields,
    extract_reliance_standard_broker_rows,
    extract_reliance_standard_schedule_a_fields,
    extract_sun_life_broker_rows,
    extract_sun_life_schedule_a_fields,
    prefer_authoritative_pomerene_fields,
    extract_ace_schedule_a_broker_rows,
    extract_ace_schedule_a_fields,
    extract_cigna_g2050a_schedule_a_fields,
    extract_delta_dental_schedule_a_fields,
    extract_first_unum_broker_rows,
    extract_first_unum_schedule_a_fields,
    extract_mount_sinai_schedule_a_fields,
    extract_metlife_bay_bridge_broker_rows,
    extract_metlife_bay_bridge_schedule_a_fields,
    extract_metlife_bay_bridge_schedule_a_summaries,
    extract_nyl_annual_policy_fields,
    extract_nyl_paid_premium_workbook,
    extract_colonial_life_schedule_a_fields,
    extract_colonial_life_broker_rows,
    extract_transamerica_schedule_a_fields,
    extract_transamerica_broker_rows,
    extract_combined_chubb_schedule_a_fields,
    extract_combined_chubb_broker_rows,
    extract_john_hancock_schedule_a_fields,
    extract_john_hancock_broker_rows,
    extract_metlife_standard_schedule_a_fields,
    extract_metlife_standard_broker_rows,
    extract_bcbsma_commission_breakdown_broker_rows,
    extract_bcbsma_schedule_a_worksheet_fields,
    extract_bcbsma_schedule_a_worksheet_summaries,
    extract_prudential_broker_rows,
    extract_prudential_schedule_a_fields,
    extract_prudential_schedule_a_summaries,
    extract_principal_short_form_broker_rows,
    extract_principal_short_form_schedule_a_fields,
    restore_principal_short_form_fields,
    extract_aetna_attached_listing_fields,
    extract_aetna_schedule_a_support_statement_fields,
    extract_litera_aetna_schedule_a_broker_rows,
    extract_litera_aetna_schedule_a_fields,
    extract_litera_lincoln_schedule_a_broker_rows,
    extract_litera_lincoln_schedule_a_fields,
    extract_position_aware_schedule_a_broker_rows,
    extract_position_aware_schedule_a_fields,
    extract_schedule_a_broker_rows,
    extract_schedule_a_fields_from_rule_labels,
    extract_summary_table_broker_rows,
    extract_standard_broker_rows,
    extract_standard_short_form_broker_rows,
    extract_standard_short_form_schedule_a_fields,
    extract_standard_schedule_a_fields,
    extract_standard_schedule_a_records,
    extract_standard_schedule_a_summaries,
    extract_united_omaha_broker_rows,
    extract_united_omaha_combined_broker_rows,
    extract_united_omaha_schedule_a_fields,
    extract_united_omaha_schedule_a_records,
    extract_united_omaha_schedule_a_summaries,
    extract_unitedhealthcare_broker_rows,
    build_groundx_schema_query,
    extract_fields_from_document_text,
    is_obvious_template_placeholder,
    is_unfilled_schedule_a_template,
    merge_schedule_a_fields,
    merge_schedule_a_broker_rows,
    remove_inapplicable_experience_rated_fields,
    money_value,
    local_schedule_a_pdf_result,
    parse_schedule_a_text,
    schedule_a_broker_compensation_fields,
    select_best_schedule_a_fields,
    supplement_schedule_a_result_with_local,
)
from app.services.schedule_a_extraction_pipeline import resolve_schedule_a_result
from app.services.field_rules import DEFAULT_FIELD_RULES
from app.services.ftwilliams_review import FTWilliamsReviewService
from app.services.mapping import map_extraction_to_rules
from app.services.schedule_a_classification import classify_schedule_a_fields


class ScheduleAExtractionTests(unittest.TestCase):
    def test_cigna_g2050a_layout_uses_certified_totals_not_table_fragments(self):
        pages = [(1, """G2050A
NON-EXPERIENCE - RATED CONTRACTS CIGNA HEALTH AND LIFE INSURANCE COMPANY
For Policy Year beginning January 1, 2025 and ending December 31, 2025
Name of plan Apollo Management, L.P.
CIGNA HEALTH AND LIFE INSURANCE COMPANY 04094A
56 Total Covered, 17 Employees 1/1/2025 12/31/2025
Dental Coverage $38,560
Evacuation Coverage $3,111
I EAP Coverage $651
Medical Coverage $845,754
Contract or Identification Number 04094A
NAIC COMPANY CODE: 67369
EMPLOYER IDENTIFICATION NUMBER: 591031071
(a) Total premiums or subscriptions charges paid to carrier $888,075
""")]
        values = {field.field_name: field.value for field in extract_cigna_g2050a_schedule_a_fields(pages)}

        self.assertEqual(values["1a. Name of Insurance Company"], "CIGNA HEALTH AND LIFE INSURANCE COMPANY")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "59-1031071")
        self.assertEqual(values["1c. NAIC Code"], "67369")
        self.assertEqual(values["1d. Contract/Policy Number"], "04094A")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "56")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "888,075")
        self.assertEqual(values["3b. Amount of Commissions"], "0")
        self.assertEqual(values["3c. Amount of Fees"], "0")

    def test_delta_dental_layout_extracts_zero_compensation_and_premium(self):
        pages = [(1, """INSURANCE INFORMATION
Delta Dental of Iowa
EIN: 42-0959302
NAIC CODE: 55786
Contract or ID Number: 43173
Approximate Number of Persons Covered @ End of Policy or Contract Year: 1980
Policy or Contract Year 01/01/2025 - 12/31/2025
Total Amount of Commissions Paid: $0.00
Total Fees Paid/Amount: $0.00
Total premiums or subscription charges paid to carrier $2,113,696.44
""")]
        values = {field.field_name: field.value for field in extract_delta_dental_schedule_a_fields(pages)}

        self.assertEqual(values["1a. Name of Insurance Company"], "DELTA DENTAL OF IOWA")
        self.assertEqual(values["1d. Contract/Policy Number"], "43173")
        self.assertEqual(values["3b. Amount of Commissions"], "0.00")
        self.assertEqual(values["3c. Amount of Fees"], "0.00")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "2,113,696.44")

    def test_first_unum_layout_normalizes_exclusive_anniversary_and_broker(self):
        pages = [(1, """INSURANCE DATA FOR SCHEDULE A (FORM 5500)
First Unum Life Insurance Company
TAX ID: 131898173 NAIC: 64297
CONTRACT NUMBER: 612696
APPROXIMATE NUMBER OF PERSONS COVERED AT END OF POLICY YEAR: 1505
DATE FOR PERIOD: FROM 01-01-2025 TO 01-01-2026
NFP Corporate Services (NY) LLC 31,892.89 .00 .00
PO Box 9101
Plainview NY 11803
NON-PARTICIPATING CONTRACTS (PREMIUMS)
TOTAL PREMIUM OR SUBSCRIPTION CHARGES PAID TO CARRIER $637,857.78
""")]
        values = {field.field_name: field.value for field in extract_first_unum_schedule_a_fields(pages)}
        rows = extract_first_unum_broker_rows(pages)

        self.assertEqual(values["1b. Insurance Carrier EIN"], "13-1898173")
        self.assertEqual(values["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "637,857.78")
        self.assertEqual(rows[0].name, "NFP Corporate Services (NY) LLC")
        self.assertEqual(rows[0].commission_total, "31,892.89")

    def test_sun_life_health_layout_does_not_turn_footnote_into_fee(self):
        pages = [(1, "Cover letter"), (2, """5500 Schedule A Insurance Information
Policy/Account Number 931755
Name of insurance carrier Sun Life and Health Insurance Company (U.S.)
EIN (Insurance Carrier) 06-0893662 NAIC code 80926
Policy or Contract Year From 01/01/2025 To 12/31/2025
1929 Approximate number of persons covered at end of policy or contract year
Total Amount of commissions paid $91,286.13
Amount of commissions paid Type of Benefit Override 1 Stop Loss Specific Only $91,286.13
Bonuses and additional payments paid 3
Bonus Amount 2 Additional Payments
Alterity Group LLC
90 Park Ave 17th Floor
New York, NY 10016
Total Premium received 01/01/2025 to 12/31/2025
Stop Loss Specific Only $2,026,913.43
Total $2,026,913.43
""")]
        values = {field.field_name: field.value for field in extract_sun_life_schedule_a_fields(pages)}
        rows = extract_sun_life_broker_rows(pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "SUN LIFE AND HEALTH INSURANCE COMPANY (U.S.)")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "06-0893662")
        self.assertEqual(values["3c. Amount of Fees"], "0.00")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "2,026,913.43")
        self.assertEqual(rows[0].name, "Alterity Group LLC")

    def test_ace_layout_keeps_missing_persons_blank_and_extracts_broker(self):
        pages = [(1, """Insurance Information: Provided to assist the Plan Administrator in completing Form 5500/Schedule A.
Contract Identification/Policy Number: N18154614
Policy Period: 06/01/2025 - 06/01/2026
Name of Insurance Company: ACE American Insurance Company
ACE American Insurance Company Tax ID Number: 95-2371728
ACE American Insurance Company NAIC Number: 22667
Approximate Number of Persons Insured: To be Provided by the Plan Administrator
Name of Agent or Broker Commissions or Fees Paid
AON CONSULTING INC
ONE LIBERTY PLAZA, 165 BROADWAY SUITE 3201
NEW YORK, NY 10006
$36,935.20
Total Premium Paid to ACE American Insurance Company: $184,676.00
""")]
        values = {field.field_name: field.value for field in extract_ace_schedule_a_fields(pages)}
        rows = extract_ace_schedule_a_broker_rows(pages)

        self.assertNotIn("1e. Persons Covered (End of Policy Year)", values)
        self.assertEqual(values["1g. Policy Year Ending Date"], "05/31/2026")
        self.assertEqual(values["3b. Amount of Commissions"], "36,935.20")
        self.assertEqual(rows[0].name, "AON CONSULTING INC")
        self.assertEqual(rows[0].commission_total, "36,935.20")

    def test_mount_sinai_invoice_layout_sums_membership_fees(self):
        pages = [(1, """2025 Apollo Management Holdings, L.P. Schedule A - Form 5500
Service Period 1/1/2025 - 1/31/2025 2/1/2025 - 2/28/2025 12/1/2025 - 12/31/2025
Number of Eligible Employees 1430 1445 1422 1433 1440 1463 1482 1526 1589 1613 1623 0
Virtual Health Center at Hudson Yards and Navigation Services Membership Fee
20,600$ 20,600$ 20,600$ 20,600$ 21,218$ 21,218$ 21,218$ 21,218$ 21,218$ 21,218$ 21,218$ - $
Relationship between Apollo Management Holdings, L.P. and Mount Sinai Solutions, LLC terminated effective December 1, 2025.
""")]
        values = {field.field_name: field.value for field in extract_mount_sinai_schedule_a_fields(pages)}

        self.assertEqual(values["1a. Name of Insurance Company"], "MOUNT SINAI SOLUTIONS LLC")
        self.assertEqual(values["1d. Contract/Policy Number"], "APOLLO")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "0")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "230,926")

    def test_bank_of_bartlett_transamerica_statement(self):
        pages = [(1, """Transamerica Life Insurance Company
SCHEDULE 'A' INFORMATION FOR SECTION 125
FOR EMPLOYER ER00000636
PLAN YEAR 01/01/2025 - 12/31/2025
NAIC 86231 Tax ID 39-0989781
BANK OF BARTLETT CANCER TLIC 24.00
6281 STAGE RD
NAIC 86231 Tax ID 39-0989781
BARTLETT TN 38134
BANK OF BARTLETT CANCER 24.00
15 5,843.28 165.60
""")]
        values = {field.field_name: field.value for field in extract_transamerica_schedule_a_fields(pages)}
        self.assertEqual(values["1d. Contract/Policy Number"], "ER00000636")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "5,843.28")
        self.assertEqual(values["3b. Amount of Commissions"], "165.60")
        rows = extract_transamerica_broker_rows(pages)
        self.assertEqual((rows[0].name, rows[0].commission_total, rows[0].city), ("BANK OF BARTLETT", "24.00", "BARTLETT"))

    def test_bank_of_bartlett_combined_chubb_statement(self):
        pages = [(1, """Combined Insurance, A CHUBB Company
5500 Annual Report Schedule A Information
1. Plan Year: 01/1/2025 – 01/31/2025
4. Insurance Company: Combined Insurance Company of America
5. Tax ID: 36-2136262
6. NAIC Code: 62146
7. Contract or ID No.: ACC; Life; – Group# 901937294; 901937295
8. Commissions Paid: AON Consulting INC: $60.39
Elizabeth Blair: $103.47
Comprehensive Wealth MGMT: $1,199.32
Charles Summers: $1,598.57
9. Basis of Premium Rates: Master Policy on File
10. Total Premium Paid: $34,100.37
11. Number of Participants: 99
""")]
        values = {field.field_name: field.value for field in extract_combined_chubb_schedule_a_fields(pages)}
        self.assertEqual(values["1d. Contract/Policy Number"], "901937294; 901937295")
        self.assertEqual(values["1g. Policy Year Ending Date"], "01/31/2025")
        self.assertEqual(values["3b. Amount of Commissions"], "2,961.75")
        self.assertEqual(len(extract_combined_chubb_broker_rows(pages)), 4)

    def test_bank_of_bartlett_john_hancock_workbook_text(self):
        pages = [(1, """Schedule A (Form 5500) Data for Year 2025
Name of Group/Group Number: Bank of Bartlett/90018
EIN Number: 01-0233346
NAIC Number: 65838
Policy Number: 30460
Inforce Count: 26
Policy Year: 01/01/2025-12/31/2025
Payee Information: William Billingsley
Payee Address: Comprehensive Wealth Mgmt
1910 Exeter Rd. Ste 2
Germantown
TN 38138
Commission Paid: $313.31
Total- Premium Paid: $23,650.80
John Hancock
""")]
        values = {field.field_name: field.value for field in extract_john_hancock_schedule_a_fields(pages)}
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "26")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "23,650.80")
        rows = extract_john_hancock_broker_rows(pages)
        self.assertEqual((rows[0].name, rows[0].commission_total, rows[0].city), ("William Billingsley", "313.31", "Germantown"))

    def test_bank_of_bartlett_standard_metlife_ocr(self):
        pages = [(3, """SCHEDULE A (Form 5500) Insurance Information
METROPOLITAN LIFE INSURANCE COMPANY
13-5581829| 65978 TM05941745 275 01/01/2025 |12/31/2025
Total amount of commissions paid Total Fees Paid / amount
22,943 0
"""), (4, """Name: PATRICK HOFFMAN
Address: 1910 EXETER RD STE 2 City: . .
GERMANTOWN ST: TN ZIP: 38138-2971
Commissions Paid Fees Paid Organization code
LIFE 12,324 | Base Commissions 03
Dental 5,567 | Base Commissions
Long Term 4,460 | Base Commissions
Disability
AD&D 592 | Base Commissions
22,943 | Sub-total 0 Sub-total
"""), (5, """Total premiums or subscription charges paid to carrier. 172,421""")]
        values = {field.field_name: field.value for field in extract_metlife_standard_schedule_a_fields(pages)}
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "275")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "172,421")
        self.assertEqual(values["3c. Amount of Fees"], "0")
        rows = extract_metlife_standard_broker_rows(pages)
        self.assertEqual((rows[0].name, rows[0].commission_total, len(rows[0].commission_rows)), ("PATRICK HOFFMAN", "22,943", 4))

    def test_local_ocr_pages_override_partial_native_text_and_ai_premium(self):
        service = ExtractionService()
        ai_result = NormalizedExtractionResult(
            provider="remote",
            fields=[
                NormalizedExtractionField(
                    field_name="10a. Total premiums or subscription charges paid to carrier",
                    value="2,440.75",
                    confidence=0.9,
                )
            ],
            raw={
                "local_ocr_pages": [
                    {
                        "page": 1,
                        "text": """
                            SCHEDULE A EARNINGS REPORT
                            AFLAC ACCOUNT # NSU61
                            NAME OF INSURANCE CARRIER
                            American Family Life Assurance Company Of New York
                            APPROXIMATE NUMBER OF PERSONS COVERED AT END OF PLAN YEAR 16
                            PLAN YEAR FROM TO 01/01/2025 - 12/31/2025
                            CONTRACT NUMBER 52-0807803
                            NAIC CODE 60380
                            TOTAL PREMIUM COLLECTED $17,794.78
                        """,
                    }
                ]
            },
        )
        with (
            patch("app.services.extractor.extract_document_text_pages", return_value=[(1, "*NSU61")]),
            patch.object(service, "_extract_schedule_a_unresolved", AsyncMock(return_value=ai_result)),
        ):
            result = asyncio.run(service.extract_schedule_a(b"%PDF-scanned", "NSU61.pdf"))

        values = {field.field_name: field.value for field in result.fields}
        self.assertEqual(
            values["10a. Total premiums or subscription charges paid to carrier"],
            "17,794.78",
        )

    def test_image_only_aflac_uses_bounded_local_ocr_when_groundx_is_unavailable(self):
        ocr_pages = [(1, """
            SCHEDULE A EARNINGS REPORT
            AFLAC ACCOUNT # NSU61
            NAME OF INSURANCE CARRIER
            American Family Life Assurance Company Of New York
            APPROXIMATE NUMBER OF PERSONS COVERED AT END OF PLAN YEAR 16
            PLAN YEAR FROM TO 01/01/2025 - 12/31/2025
            Contract Number 52-0807803
            NAIC CODE 60380
            TOTAL PREMIUM COLLECTED $17,794.78
            INSURANCE FEES AND COMMISSIONS PAID TO AGENTS
            COMMISSIONS PAID FEES PAID
            JENNIFER LUBELSKY $661.32 $30.34
        """)]
        with (
            patch("app.services.extractor.extract_pdf_text_pages", return_value=[(1, "")]),
            patch("app.services.extractor.extract_pdf_layout_text_pages", return_value=[(1, "")]),
            patch("app.services.extractor.extract_image_only_pdf_ocr_pages", return_value=ocr_pages),
            patch("app.services.extractor.extract_schedule_a_worksheet_summaries_from_pdf_text", return_value=[]),
        ):
            result = local_schedule_a_pdf_result(b"%PDF-scanned", "NSU61.pdf")

        values = {field.field_name: field.value for field in result.fields}
        self.assertEqual(values["1a. Name of Insurance Company"], "American Family Life Assurance Company Of New York")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "52-0807803")
        self.assertEqual(values["1d. Contract/Policy Number"], "NSU61")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "17,794.78")
        self.assertEqual(len(result.schedule_a_broker_rows), 1)
        self.assertEqual(result.schedule_a_broker_rows[0].commission_total, "661.32")
        self.assertEqual(result.raw["source"], "local_pdf_ocr_parser")

    def test_aflac_ocr_rows_strip_artifacts_without_dropping_brokers(self):
        pages = [(1, """
            SCHEDULE A EARNINGS REPORT
            AFLAC ACCOUNT # NSU61
            COMMISSIONS PAID FEES PAID
            — JENNIFER LUBELSKY $661.32 $30.34
            Ss JASON M GREB $644.31 $12.33
            —I WORLD INSURANCE ASSOCIATES LLC $317.15 $0.00
            OO CHRISTOPHER J ROTH $221.80 $0.00
            GROE INC $47.60 $1.52
            OO GRAND TOTAL $1,892.20 $44.34
        """)]

        rows = extract_aflac_broker_rows(pages)

        self.assertEqual(
            [row.name for row in rows],
            [
                "JENNIFER LUBELSKY",
                "JASON M GREB",
                "WORLD INSURANCE ASSOCIATES LLC",
                "CHRISTOPHER J ROTH",
                "GROF INC",
            ],
        )
        self.assertEqual(rows[0].fee_total, "30.34")
        self.assertEqual(rows[-1].commission_total, "47.60")

    def test_aflac_ocr_premium_allows_scanner_rule_prefixes(self):
        pages = [(1, """
            — SCHEDULE A EARNINGS REPORT AFLAC
            — AFLAC ACCOUNT # NSU79
            — Contract Number 52-0807803 NAIC CODE 60380
            — TOTAL PREMIUM COLLECTED
            — $315.12
            — INSURANCE FEES AND COMMISSIONS PAID TO AGENTS
        """)]

        values = {field.field_name: field.value for field in extract_aflac_schedule_a_fields(pages)}

        self.assertEqual(
            values["10a. Total premiums or subscription charges paid to carrier"],
            "315.12",
        )

    def test_colonial_life_ocr_layout_extracts_identity_totals_and_brokers(self):
        pages = [(2, """
            Insurance Data for Schedule A Form 5500
            Name of Carrier: The Paul Revere Life Insurance Company
            Carrier EIN: 04-1590994
            Carrier NAIC Code: 67598
            Billing Control Number: E4020418
            Plan Year Date Range: 05/01/2025 - 04/30/2026
            Organization Code For Agents/Producers: 3
            Total Paid Premium: $1,722.37
            APPROXIMATE NUMBER OF PERSONS COVERED IN APRIL 2026: 2
            Julie Ann Klimchak $0.00 $15.90 $15.90 $0.00
            Jnaz Inc $0.00 $3.81 $3.81 $0.00
            Grand Totals $0.00 $19.71 $19.71 $0.00
        """)]
        values = {field.field_name: field.value for field in extract_colonial_life_schedule_a_fields(pages)}
        brokers = extract_colonial_life_broker_rows(pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "The Paul Revere Life Insurance Company")
        self.assertEqual(values["1d. Contract/Policy Number"], "E4020418")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "1,722.37")
        self.assertEqual(values["3b. Amount of Commissions"], "19.71")
        self.assertEqual(values["3c. Amount of Fees"], "0.00")
        self.assertEqual([row.name for row in brokers], ["JULIE ANN KLIMCHAK", "JNAZ INC"])
        self.assertTrue(all(row.fee_total == "0.00" for row in brokers))

    def test_colonial_ocr_prefers_bcn_and_final_person_count(self):
        pages = [(1, """
            Insurance Data for Schedule A Form 5500
            THE PAUL REVERE LIFE INSURANCE COMPANY
            Name of Carrier: The Paul Revere Life Insurance Company
            Carrier EIN: 04-1590994
            Carrier NAIC Code: 67598
            BCN: E4020418
            Billing Control Number: 4020418
            Plan Year Date Range: 05/01/2025 - 04/30/2026
            APPROXIMATE NUMBER OF PERSONS COVERED IN APRIL 2026: 2
        """)]

        values = {field.field_name: field.value for field in extract_colonial_life_schedule_a_fields(pages)}

        self.assertEqual(values["1d. Contract/Policy Number"], "E4020418")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "2")

    def test_colonial_labelled_person_count_replaces_flattened_ocr_guess(self):
        fields = [
            NormalizedExtractionField(
                field_name="1e. Persons Covered (End of Policy Year)",
                value="2026",
                confidence=0.9,
            )
        ]
        pages = [(2, """
            Insurance Data for Schedule A Form 5500
            THE PAUL REVERE LIFE INSURANCE COMPANY
            Name of Carrier: The Paul Revere Life Insurance Company
            Carrier EIN: 04-1590994
            Carrier NAIC Code: 67598
            Billing Control Number: E4020418
            Plan Year Date Range: 05/01/2025 - 04/30/2026
            APPROXIMATE NUMBER OF PERSONS COVERED IN APRIL 2026: 2
            Total Paid Premium: $1,722.37
        """)]

        corrected = prefer_authoritative_colonial_fields(fields, pages)
        values = {field.field_name: field.value for field in corrected}

        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "2")

    def test_nyl_short_year_workbook_sums_policy_and_broker_transactions(self):
        from openpyxl import Workbook

        workbook = Workbook()
        premiums = workbook.active
        premiums.title = "PaidPremiumData"
        premiums.append(["ClientName", "PayorId", "Underwriter", "PolicyNumber", "PremPeriod", "Product_Benefit", "ApplyDate", "AppliedAmt"])
        premiums.append(["BWD", "1", "CLICNY", "SGN0600973", datetime(2025, 6, 1), "Life - Basic", datetime(2025, 6, 30), 100])
        premiums.append(["BWD", "1", "CLICNY", "SGN0600973", datetime(2026, 4, 1), "Life - Voluntary", datetime(2026, 4, 30), 50])
        premiums.append(["BWD", "1", "CLICNY", "VDY0600189", datetime(2026, 4, 1), "STD - Voluntary", datetime(2026, 4, 30), 25])
        commissions = workbook.create_sheet("Commissions")
        commissions.append(["Broker Number", "Broker Name", "Payment Issuance date", "Commission Type", "Client Name", "Policy Number", "Premium Start Date", "Commission Amount"])
        commissions.append(["GPO-1", "BENEFITMALL", datetime(2025, 8, 1), "GAFE", "BWD", "SGN0600973", datetime(2025, 6, 1), 12.5])
        commissions.append(["GPO-2", "WORLD INSURANCE ASSOCIATES LLC", datetime(2025, 8, 1), "GAFE", "BWD", "VDY0600189", datetime(2025, 6, 1), 7.5])
        buffer = BytesIO()
        workbook.save(buffer)

        result = extract_nyl_paid_premium_workbook(buffer.getvalue(), "BWD NYL.xlsx")
        self.assertIsNotNone(result)
        values = {field.field_name: field.value for field in result.fields}
        self.assertEqual(values["1a. Name of Insurance Company"], "New York Life Group Insurance Company of NY")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "13-2556568")
        self.assertEqual(values["1c. NAIC Code"], "64548")
        self.assertEqual(values["1d. Contract/Policy Number"], "SGN0600973; VDY0600189")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "06/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "04/30/2026")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "175")
        self.assertEqual(values["3b. Amount of Commissions"], "20")
        self.assertEqual(len(result.schedule_a_worksheet_summaries), 2)
        self.assertEqual(len(result.schedule_a_broker_rows), 2)
        self.assertTrue(all(row.fee_total == "0" for row in result.schedule_a_broker_rows))

    def test_pomerene_eyemed_keeps_two_policies_separate(self):
        pages = [(1, """
            Vision Insurance Information For Form 5500
            Information Compiled By: EyeMed Vision Care on behalf of the Fidelity Security Life Insurance Company
            Report Start Date Report End Date
            1/1/2025 12/31/2025
            Name of Plan Contract or ID # Enrollment Group subscribers covered subscribers and dependents covered EIN NAIC Amount
            POMERENE HOSPITAL 10049071001 POMERENE HOSPITAL 81 172 430949844 71870 $2,451.00
            POMERENE HOSPITAL BUY UP 10049061001 POMERENE HOSPITAL BUY UP 174 432 430949844 71870 $33,717.70
            Total: $36,168.70
            Payee Name Contract or ID # Address Line 1 City State Zip Code Amount
            Hummel Group 10049061001 461 Wadsworth Road PO Box 3 Orrville OH 44667 $3,406.53
            Hummel Group 10049071001 461 Wadsworth Road PO Box 3 Orrville OH 44667 $245.62
        """)]

        summaries = extract_eyemed_schedule_a_summaries(pages)

        self.assertEqual([summary.account_number for summary in summaries], ["10049071001", "10049061001"])
        values = [{value.label: value.value for value in summary.values} for summary in summaries]
        self.assertEqual(values[0]["Persons covered"], "172")
        self.assertEqual(values[0]["Total nonexperience premium"], "2,451.00")
        self.assertEqual(values[0]["Broker payment total"], "245.62")
        self.assertEqual(values[1]["Persons covered"], "432")
        self.assertEqual(values[1]["Total nonexperience premium"], "33,717.70")
        self.assertEqual(values[1]["Broker payment total"], "3,406.53")
        brokers = extract_eyemed_broker_rows(pages)
        self.assertEqual(len(brokers), 2)
        self.assertEqual(brokers[0].address_line_1, "461 Wadsworth Road")
        self.assertEqual(brokers[0].address_line_2, "PO Box 3")
        self.assertEqual(brokers[0].city, "Orrville")

    def test_pomerene_guardian_portal_statement_extracts_policy_totals_and_brokers(self):
        pages = [(1, """
            Guardian 2025 Schedule A/5500 Information
            From 01/01/2025 To 12/31/2025
            Plan Name POMERENE HOSPITAL Plan Number 00579175
            Guardian's EIN 13-5123390 Guardian's NAIC 64246
            Approximate number of employees covered at the end of the plan year 418
            0009L838 REDTAIL LTD $6,431.27
            9999 BREWSTER PLACE #100 POWELL OH 43065
            000K2652 HUMMEL GROUP INC $6,431.27
            PO BOX 250 BERLIN OH 44610 111
            000NM733 HUMMEL GROUP INC $0.00
            Total commissions for plan $12,862.54
            Total Fees Paid $0.00
            Total premium paid $98,942.65
        """)]

        fields = {field.field_name: field.value for field in extract_guardian_schedule_a_fields(pages)}
        brokers = extract_guardian_broker_rows(pages)

        self.assertEqual(fields["1d. Contract/Policy Number"], "00579175")
        self.assertEqual(fields["1e. Persons Covered (End of Policy Year)"], "418")
        self.assertEqual(fields["3b. Amount of Commissions"], "12,862.54")
        self.assertEqual(fields["3c. Amount of Fees"], "0.00")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "98,942.65")
        self.assertEqual([(row.name, row.commission_total) for row in brokers], [("REDTAIL LTD", "6,431.27"), ("HUMMEL GROUP INC", "6,431.27")])

    def test_guardian_fee_with_blank_recipient_uses_matching_contract_broker(self):
        pages = [(2, """
            Guardian Life Insurance Company of America
            Plan Number : 00398015 EIN : 13-5123390 NAIC: 64246
            Data for Period From : 1/1/25 To : 12/31/25
            The following figure represents commissions that are to be reported on Schedule A, Line 3, Element (b):
            Contract Identification Name and Address of Recipient of Commissions
            0002Z407 GALLAGHER BENEFIT SERVICES INC
            PARK CENTRAL 7/ 12750 MERIT DR SUITE 1000 DALLAS TX 7525
            Dental (Insured) 10,483.70
            Total For Contract: 10,483.70
            Total Commissions Paid On Plan: 10,483.70
            The following figure represents fees that are to be reported on Schedule A, Line 3, Element (c):
            Contract Identification Name of Recipient of Fees Amount
            0002Z407 $10,713.11
            Total Fees Paid $10,713.11
            Totals: 349,456.83
        """)]

        brokers = extract_guardian_broker_rows(pages)

        self.assertEqual(len(brokers), 1)
        self.assertEqual(brokers[0].name, "GALLAGHER BENEFIT SERVICES INC")
        self.assertEqual(brokers[0].commission_total, "10,483.70")
        self.assertEqual(brokers[0].fee_total, "10,713.11")
        self.assertEqual(brokers[0].fee_rows[0].amount, "10,713.11")
        self.assertEqual(brokers[0].fee_rows[0].purpose, "FEES")

    def test_pomerene_guardian_letter_uses_plan_number_not_broker_code(self):
        pages = [(1, """
            Guardian Life Insurance Company of America
            Plan Number : 00579270 EIN : 13-5123390 NAIC: 64246
            Name of Plan : POMERENE HOSPITAL
            Data for Period From : 1/1/25 To : 12/31/25
            approximate number of employees covered at the end of the plan year : 63
            000K215 LIFETIME FINANCIAL GROWTH OF N Short Term Disability 38.48 Total For Contract: 38.48
            000K2652 HUMMEL GROUP INC PO BOX 250 BERLIN OH 44610 111 Short Term Disability 4,735.73 Total For Contract: 4,735.73
            Total Commissions Paid On Plan: 4,774.21
            000NM733 HUMMEL GROUP INC $0.00 Total Fees Paid $0.00
            Gross Premium Paid Short Term Disability 29,598.29 Totals: 29,598.29
        """)]

        fields = {field.field_name: field.value for field in extract_guardian_schedule_a_fields(pages)}

        self.assertEqual(fields["1d. Contract/Policy Number"], "00579270")
        self.assertEqual(fields["3b. Amount of Commissions"], "4,774.21")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "29,598.29")

    def test_pomerene_american_heritage_aggregates_same_contract_benefits(self):
        pages = [(1, """
            POMERENE HOSPITAL(MH301)
            Plan/Contract Year:1/1/2025-12/31/2025
            Part I (a) Name of Insurance Carrier: American Heritage Life Insurance Company
            Part I (b) EIN: 59-0781901 Part I (c) NAIC Code: 60534
            Accident Account: MH301 Accident 83 $31,169.61 $9,726.44 $0.00
            CGI VOLUNTARY BENEFITS INC 8HRK0 $121.20 $0.00
            HUMMEL GROUP INC 8Y0K0 $6,652.62 $0.00
        """), (3, """
            POMERENE HOSPITAL(MH301)
            Critical Illness Account: MH301 Critical Illness 31 $13,368.95 $2,517.41 $0.00
            CGI VOLUNTARY BENEFITS INC 8HRK0 $40.78 $0.00
            HUMMEL GROUP INC 8Y0K0 $1,347.88 $0.00
        """), (5, """
            POMERENE HOSPITAL(MH301)
            Universal Life Account: MH301 Universal Life 71 $21,166.88 $650.77 $0.00
            GALLAGHER BENEFIT SVCS INC 0HW70 $253.91 $0.00
        """), (6, "Grand Total $65,705.44 $12,894.62 $0.00")]

        fields = {field.field_name: field.value for field in extract_american_heritage_schedule_a_fields(pages)}
        brokers = extract_american_heritage_broker_rows(pages)

        self.assertEqual(fields["1d. Contract/Policy Number"], "MH301")
        self.assertEqual(fields["1e. Persons Covered (End of Policy Year)"], "83")
        self.assertEqual(fields["3b. Amount of Commissions"], "12,894.62")
        self.assertEqual(fields["3c. Amount of Fees"], "0.00")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "65,705.44")
        self.assertEqual(next(row for row in brokers if row.name == "HUMMEL GROUP INC").commission_total, "8,000.50")

    def test_pomerene_sun_life_extracts_complete_schedule_a(self):
        pages = [(2, """
            5500 Schedule A Insurance Information
            Policy/Account Number 924948
            Name of insurance carrier Sun Life Assurance Company of Canada EIN 38-1082080 NAIC code 80802
            Policy or Contract Year From 01/01/2025 To 12/31/2025
            Approximate number of persons covered at end of policy or contract year 62
            Total Amount of commissions paid $6,330.49
            Additional payments paid $0.00
            Gallagher Benefit Services Inc 2850 Golf Rd 5th Fl Rolling Meadows, IL 60008 Organization Code 3
            Total Premium received 01/01/2025 to 12/31/2025 Total $32,478.18
        """)]

        fields = {field.field_name: field.value for field in extract_sun_life_schedule_a_fields(pages)}
        brokers = extract_sun_life_broker_rows(pages)

        self.assertEqual(fields["1d. Contract/Policy Number"], "924948")
        self.assertEqual(fields["1e. Persons Covered (End of Policy Year)"], "62")
        self.assertEqual(fields["3b. Amount of Commissions"], "6,330.49")
        self.assertEqual(fields["3c. Amount of Fees"], "0.00")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "32,478.18")
        self.assertEqual(brokers[0].name, "Gallagher Benefit Services Inc")

        resolved = resolve_schedule_a_result(
            NormalizedExtractionResult(
                provider="Local PDF parser",
                fields=extract_sun_life_schedule_a_fields(pages),
                schedule_a_broker_rows=extract_sun_life_broker_rows(pages),
            )
        )
        self.assertEqual(resolved.raw["extraction_quality"]["error_count"], 0)
        self.assertEqual(resolved.raw["extraction_quality"]["cross_field_errors"], [])
        self.assertTrue(all(field.decision == "AUTOMATIC" for field in resolved.fields))
        self.assertTrue(all(row.decision == "AUTOMATIC" for row in resolved.schedule_a_broker_rows))

    def test_pomerene_sun_life_table_header_is_not_used_as_carrier_name(self):
        pages = [(2, """
            5500 Schedule A Insurance Information
            Policy/Account Number 924948
            Name of insurance carrier
            EIN (Insurance Carrier) NAIC Code From To
            Sun Life Assurance Company of Canada 38-1082080 80802 01/01/2025 12/31/2025
            Policy or Contract Year From 01/01/2025 To 12/31/2025
            62 Approximate number of persons covered at end of policy or contract year
            Total Amount of commissions paid $6,330.49
            Total Premium received 01/01/2025 to 12/31/2025 Total $32,478.18
        """)]

        fields = {field.field_name: field.value for field in extract_sun_life_schedule_a_fields(pages)}

        self.assertEqual(
            fields["1a. Name of Insurance Company"],
            "SUN LIFE ASSURANCE COMPANY OF CANADA",
        )

    def test_pomerene_reliance_ocr_text_extracts_complete_schedule_a(self):
        pages = [(1, """
            reliance standard INSURANCE INFORMATION FORM 5500 - SCHEDULE A
            EIN: 36-0883760 NAIC: 68381 ORG. NUMBER: 3
            Policyholder Name: Pomerene Hospital Policy Number: GL160111
            Policy Type: GROUP LIFE AND ACCIDENTAL DEATH AND DISMEMBERMENT
            Number of covered lives: Beginning: 351 Ending: 319
            Policy Contract Year: 01/01/2025 to 12/31/2025
            Total Premium: $10,878.32
            Payee Name: Gallagher Benefit Services Inc
            Payee Address: Mail Stop: 072103 P. O. Box 4135 Clinton, IA 52732
            Total Commission: $1,522.94
            Total Administrative and Other Fees: $384.40
        """)]

        fields = {field.field_name: field.value for field in extract_reliance_standard_schedule_a_fields(pages)}
        brokers = extract_reliance_standard_broker_rows(pages)

        self.assertEqual(fields["1d. Contract/Policy Number"], "GL160111")
        self.assertEqual(fields["1e. Persons Covered (End of Policy Year)"], "319")
        self.assertEqual(fields["3b. Amount of Commissions"], "1,522.94")
        self.assertEqual(fields["3c. Amount of Fees"], "384.40")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "10,878.32")
        self.assertEqual(brokers[0].name, "Gallagher Benefit Services Inc")

    def test_pomerene_reliance_vertical_labels_do_not_replace_semantic_policy(self):
        pages = [(1, """
            reliance standard INSURANCE INFORMATION FORM 5500 - SCHEDULE A
            EIN: 36-0883760
            NAIC: 68381
            ORG. NUMBER: 3
            Policyholder Name:
            Policy Number:
            Policy Type:
            Number of covered lives:
            Policy Contract Year:
            Pomerene Hospital
            GL160111
            GROUP LIFE AND ACCIDENTAL DEATH AND DISMEMBERMENT
            Beginning: 351
            01/01/2025 to 12/31/2025
            Ending: 319
            Total Premium: $10,878.32
            Payee Name : Gallagher Benefit Services Inc    Total Administrative and Other Fees : $384.40
            Payee Address : Mail Stop : 072103    Total Amount of Compensation : $1,907.34
            P. O. Box 4135 Clinton, IA 52732
            Total Commission: $1,522.94
        """)]
        provider_fields = [
            NormalizedExtractionField(
                field_name="1d. Contract/Policy Number",
                value="GL160111",
                confidence=0.97,
                page=1,
                source_text="EyeLevel semantic policy record",
            )
        ]

        merged = prefer_authoritative_pomerene_fields(provider_fields, pages)
        values = {field.field_name: field.value for field in merged}
        brokers = extract_reliance_standard_broker_rows(pages)

        self.assertEqual(values["1d. Contract/Policy Number"], "GL160111")
        self.assertEqual(brokers[0].name, "Gallagher Benefit Services Inc")
        self.assertEqual(brokers[0].address_line_1, "Mail Stop : 072103")
        self.assertEqual(brokers[0].address_line_2, "P. O. Box 4135")
        self.assertEqual(brokers[0].city, "Clinton")
        self.assertEqual(brokers[0].state, "IA")
        self.assertEqual(brokers[0].zip_code, "52732")

    def test_pomerene_reliance_drops_money_only_provider_broker_name(self):
        pages = [(1, """
            reliance standard INSURANCE INFORMATION FORM 5500 - SCHEDULE A
            EIN: 36-0883760 NAIC: 68381 ORG. NUMBER: 3
            Policyholder Name: Pomerene Hospital Policy Number: GL160111
            Number of covered lives: Beginning: 351 Ending: 319
            Policy Contract Year: 01/01/2025 to 12/31/2025
            Total Premium: $10,878.32
            Total Commission: $1,522.94
            Total Administrative and Other Fees: $384.40
        """)]
        provider_fields = [
            NormalizedExtractionField(
                field_name="3a. Name of Agent/Broker/Person",
                value="384.40",
                confidence=0.99,
            )
        ]

        merged = prefer_authoritative_pomerene_fields(provider_fields, pages)
        values = {field.field_name: field.value for field in merged}

        self.assertNotIn("3a. Name of Agent/Broker/Person", values)
        self.assertEqual(values["3b. Amount of Commissions"], "1,522.94")
        self.assertEqual(values["3c. Amount of Fees"], "384.40")
    def test_aig_welfare_plan_extracts_authoritative_identity_and_zero_broker_row(self):
        pages = [
            (
                1,
                """
                AIG INFORMATION NECESSARY TO COMPLETE SCHEDULE A (FORM 5500) WELFARE PLAN
                This information is provided by: AIG Property Casualty, U.S.
                Insurance Company: National Union Fire Ins. Co. of Pittsburgh, PA, EIN 25-0687550, NAIC 012-19445
                Contract Number or Identification GTP 0009118182-B
                POLICY/CONTRACT YEAR From 04/01/2025 To 03/31/2026
                Insurance Fees or Commissions Paid to General Agents or Brokers
                3(b) Name and Address of each recipient
                3(c) Amount of Commissions Paid
                3(d) Amount of Fees Paid
                3(d) Purpose for which fees paid
                Arthur J. Gallagher Risk Management Services LLC
                500 N Brand Blvd, Ste 100, Glendale, CA 91203-3018
                $0.00
                Accidental Death & Dismemberment $0.00
                """,
            )
        ]

        values = {field.field_name: field.value for field in extract_aig_schedule_a_fields(pages)}
        rows = extract_aig_broker_rows(pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "National Union Fire Ins. Co. of Pittsburgh, PA")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "25-0687550")
        self.assertEqual(values["1c. NAIC Code"], "19445")
        self.assertEqual(values["1d. Contract/Policy Number"], "GTP0009118182-B")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "04/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "03/31/2026")
        self.assertEqual(values["3b. Amount of Commissions"], "0.00")
        self.assertEqual(values["3c. Amount of Fees"], "0")
        self.assertNotIn("3a. Name of Agent/Broker/Person", values)
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "0.00")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "Arthur J. Gallagher Risk Management Services LLC")
        self.assertEqual(rows[0].address_line_1, "500 N Brand Blvd, Ste 100")
        self.assertEqual(rows[0].city, "Glendale")
        self.assertEqual(rows[0].state, "CA")
        self.assertEqual(rows[0].zip_code, "91203-3018")
        self.assertEqual(rows[0].commission_total, "0.00")
        self.assertEqual(rows[0].fee_total, "0")

        provider_fields = [
            NormalizedExtractionField(
                field_name="3a. Name of Agent/Broker/Person",
                value="Arthur J. Gallagher Risk Management Services LLC",
                confidence=0.99,
            )
        ]
        merged = prefer_authoritative_aig_fields(provider_fields, pages)
        self.assertNotIn(
            "3a. Name of Agent/Broker/Person",
            {field.field_name: field.value for field in merged},
        )

    def test_anthem_combined_report_extracts_one_complete_broker_and_total_premium(self):
        pages = [
            (
                2,
                """
                Information For Completion of ERISA 5500 Schedule A
                Name of Plan: FULLER THEOLOGICAL SEMINARY
                For Period: 04/01/2025 - 03/31/2026
                Customer ID: L05472
                Anthem Blue Cross Life and Health Insurance Company (G0360) 95-4331852 62825 HEALTH INDEMNITY
                Anthem Blue Cross Life and Health Insurance Company (G0360) 95-4331852 62825 Health PPO
                Blue Cross of California (G0200) 95-3760980 Health HMO
                GALLAGHER BENEFIT SERVICES INC - 323 WEST LAKESIDE AVENUE SUITE 410, CLEVELAND, OH 44113
                $125,287.31 $4,220.87
                """,
            ),
            (
                3,
                """
                Part III Welfare Benefit Contract Information
                Health PPO $1,547,077 99/248
                Health HMO $2,324,970 162/377
                Health Indemnity $57,114 229/420
                """,
            ),
        ]

        values = {field.field_name: field.value for field in extract_anthem_schedule_a_fields(pages)}
        rows = extract_anthem_broker_rows(pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "Anthem Blue Cross Life and Health Insurance Company")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "95-4331852")
        self.assertEqual(values["1c. NAIC Code"], "62825")
        self.assertEqual(values["1d. Contract/Policy Number"], "L05472")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "420")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "04/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "03/31/2026")
        self.assertEqual(values["3b. Amount of Commissions"], "125,287.31")
        self.assertEqual(values["3c. Amount of Fees"], "4,220.87")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "3,929,161")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "GALLAGHER BENEFIT SERVICES INC")
        self.assertEqual(rows[0].address_line_1, "323 WEST LAKESIDE AVENUE SUITE 410")
        self.assertEqual(rows[0].city, "CLEVELAND")
        self.assertEqual(rows[0].state, "OH")
        self.assertEqual(rows[0].zip_code, "44113")
        self.assertEqual(rows[0].organization_code, "3")
        self.assertEqual(rows[0].commission_total, "125,287.31")
        self.assertEqual(rows[0].fee_total, "4,220.87")

    def test_anthem_combined_report_supports_multiple_brokers_and_member_count(self):
        pages = [
            (
                2,
                """
                Information For Completion of ERISA 5500 Schedule A
                For Period : 01/01/2025 - 12/31/2025
                Customer ID : 300683
                Anthem Blue Cross (G1921) 23-7391136 55093 Health PPO
                EMERSON ROGERS
                LLC - 5200 N PALM AVE #114,
                FRESNO, CA 93704
                $0.00 $30,270.00
                RSC INS BROKERAGE
                INC - 2101 FLORENCE AVE,
                CINCINNATI, OH 45206
                $104,203.61 $0.00
                """,
            ),
            (
                3,
                """
                Part III Welfare Benefit Contract Information
                Health PPO $3,170,109 175/205
                """,
            ),
        ]

        values = {field.field_name: field.value for field in extract_anthem_schedule_a_fields(pages)}
        rows = extract_anthem_broker_rows(pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "Anthem Blue Cross")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "205")
        self.assertEqual(values["3b. Amount of Commissions"], "104,203.61")
        self.assertEqual(values["3c. Amount of Fees"], "30,270")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].name, "EMERSON ROGERS LLC")
        self.assertEqual(rows[1].name, "RSC INS BROKERAGE INC")

    def test_aflac_earnings_report_extracts_identity_and_every_broker_row(self):
        plain_pages = [
            (
                1,
                """
                SCHEDULE A EARNINGS REPORT
                Group Number Group Covered Count
                NBX36 57
                Total Premium Collected
                $44,327.27
                PLAN YEAR 01/01/2025-12/31/2025
                NAIC CODE
                60380
                Name of Insurance Carrier
                AFLAC
                CONTRACT NUMBER
                52-0807803
                """,
            )
        ]
        layout_pages = [
            (
                1,
                """
                SCHEDULE A EARNINGS REPORT
                Agent Address Block w/ Full Name       Commissions Paid       Fees Paid
                JAMES K FULATER
                9840 57TH AVE
                APT 3M
                CORONA, NY 11368
                    $4,608.16       $497.50
                - RSC INSURANCE BROKERAGE
                INC
                160 FEDERAL ST FL 2
                BOSTON, MA 02110
                    $196.16       $0.00
                """,
            ),
            (
                2,
                """
                KENNETH WELLER
                160 BEDFORD AVE APT 2R
                BROOKLYN, NY 11249
                    $23.73       $96.03
                Sum: $4,828.05 $593.53
                """,
            ),
        ]

        values = {field.field_name: field.value for field in extract_aflac_schedule_a_fields(plain_pages)}
        rows = extract_aflac_broker_rows(layout_pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "AFLAC")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "52-0807803")
        self.assertEqual(values["1c. NAIC Code"], "60380")
        self.assertEqual(values["1d. Contract/Policy Number"], "NBX36")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "57")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "44,327.27")
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1].name, "RSC INSURANCE BROKERAGE INC")
        self.assertEqual(rows[2].fee_total, "96.03")

    def test_vsp_report_uses_labelled_carrier_and_policy_year_values(self):
        pages = [
            (
                1,
                """
                VSP vision care
                SCHEDULE A (Form 5500) Insurance Information
                Group No: 30027869
                Ins. Carrier: Vision Service Plan
                Ins. Carrier NAIC Code No: 39616
                Ins. Carrier FEIN: 061227840
                Benefit Type: Vision Care
                Policy or Contract Year: 01/01/2025 - 12/31/2025
                Approximate Number of Persons Covered at the End of Policy or Contract Year: 549
                Total Payments Made to Carrier: $66,964.14
                """,
            )
        ]

        values = {field.field_name: field.value for field in extract_vsp_schedule_a_fields(pages)}

        self.assertEqual(values["1a. Name of Insurance Company"], "Vision Service Plan")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "06-1227840")
        self.assertEqual(values["1c. NAIC Code"], "39616")
        self.assertEqual(values["1d. Contract/Policy Number"], "30027869")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "549")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "66,964.14")

    def test_vsp_payment_claim_fee_summary_maps_to_experience_rated_lines(self):
        pages = [
            (
                1,
                """
                VSP vision care
                Schedule A Form (5500) Insurance Information
                Group ID: 30105423
                Insurance Carrier: Vision Service Plan
                Insurance Carrier NAIC Code: 47029
                Insurance Carrier FEIN: 222777159
                Policy or Contract Year: 01/01/2025 - 12/31/2025
                Approximate Number of Persons Covered at the End of Policy or Contract Year: 3,055
                Total Administrative Fees Paid to Carrier: $67,975.56
                Total Payments Made to Carrier: $503,524.06
                Total Claims Paid by Carrier: $424,707.02
                """,
            )
        ]

        values = {field.field_name: field.value for field in _extract_fields_from_pages(pages)}

        self.assertEqual(values["9a. Premiums: (1) Amount Received"], "503,524.06")
        self.assertEqual(values["9b(1). Benefit Charges (1) Claims paid"], "424,707.02")
        self.assertEqual(values["9c(1)(B). Administrative service or other fees"], "67,975.56")
        self.assertNotIn("3c. Amount of Fees", values)
        self.assertNotIn("10a. Total premiums or subscription charges paid to carrier", values)

    def test_vsp_full_labels_and_compact_text_override_generic_false_matches(self):
        pages = [
            (
                1,
                """
                Schedule A Form (5500) Insurance Information
                Fuller Theological Seminary BERNADETTE BARBER Group ID: 30110579135 N OAKLAND AVE
                Insurance Carrier: Vision Service PlanPASADENA CA 91101-1713
                Insurance Carrier NAIC Code: N/A
                Insurance Carrier FEIN: 941632821 Benefit Type: Vision Care
                Policy or Contract Year: 04/01/2025 - 03/31/2026
                Approximate Number of Persons Covered at the End of Policy or Contract Year: 255
                Total Payments Made to Carrier: $39,795.20
                """,
            )
        ]

        values = {field.field_name: field.value for field in _extract_fields_from_pages(pages)}

        self.assertEqual(values["1a. Name of Insurance Company"], "Vision Service Plan")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "94-1632821")
        self.assertNotIn("1c. NAIC Code", values)
        self.assertEqual(values["1d. Contract/Policy Number"], "30110579")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "255")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "04/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "03/31/2026")
        self.assertEqual(
            values["10a. Total premiums or subscription charges paid to carrier"],
            "39,795.20",
        )

    def test_vsp_compact_report_extracts_one_commission_only_broker_row(self):
        pages = [
            (
                1,
                """
                VSP vision care Schedule A Form (5500) Insurance Information
                Insurance Carrier: Vision Service Plan
                Insurance Fees and Commissions Paid to Agents and Brokers:
                Commissions/Fees Paid for Policy Agent or Broker or Contract Year
                Gallagher Benefit Services Inc $3,984.882850 Golf RdRolling Meadows IL 60008
                """,
            )
        ]

        rows = extract_vsp_broker_rows(pages)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "Gallagher Benefit Services Inc")
        self.assertEqual(rows[0].address_line_1, "2850 Golf Rd")
        self.assertEqual(rows[0].city, "Rolling Meadows")
        self.assertEqual(rows[0].state, "IL")
        self.assertEqual(rows[0].zip_code, "60008")
        self.assertEqual(rows[0].organization_code, "3")
        self.assertEqual(rows[0].commission_total, "3,984.88")
        self.assertEqual(rows[0].fee_total, "0")
        self.assertEqual(rows[0].commission_rows[0].purpose, "COMMISSIONS")
        self.assertEqual(rows[0].fee_rows, [])

        fields = {field.field_name: field.value for field in _extract_fields_from_pages(pages)}
        self.assertEqual(fields["3a. Name of Agent/Broker/Person"], "Gallagher Benefit Services Inc")
        self.assertEqual(fields["3b. Amount of Commissions"], "3,984.88")
        self.assertEqual(fields["3c. Amount of Fees"], "0")
        self.assertEqual(fields["3d. Purpose"], "COMMISSIONS")
        self.assertEqual(fields["3e. Organizational Code"], "3")

    def test_vsp_duplicate_pages_do_not_double_broker_commission(self):
        statement = """
        VSP vision care Schedule A Form (5500) Insurance Information
        Group ID: 40152233
        Insurance Carrier: Vision Service Plan
        Insurance Carrier NAIC Code: 39616
        Insurance Carrier FEIN: 061227840
        Policy or Contract Year: 01/01/2025 - 12/31/2025
        Approximate Number of Persons Covered at the End of Policy or Contract Year: 118
        Total Payments Made to Carrier: $21,711.02
        Commissions/Fees Paid for Policy Agent or Broker or Contract Year
        NFP Corporate Services NY LLC $972.09
        200 Park Ave 32nd FL
        New York NY 10166
        """
        pages = [(1, statement), (2, statement)]

        rows = extract_vsp_broker_rows(pages)
        fields = {field.field_name: field.value for field in _extract_fields_from_pages(pages)}

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].commission_total, "972.09")
        self.assertEqual(fields["3b. Amount of Commissions"], "972.09")

    def test_hartford_statement_extracts_identity_premium_lives_and_bonus_as_fee(self):
        pages = [
            (
                2,
                """
                The Hartford Group Benefits Division
                Annual Statement of Premiums and Producer Compensation
                Plan/Policy Year - 01/01/2025 to 12/31/2025
                Name of Insurance Carrier EIN NAIC Code Policy Number
                HARTFORD LIFE AND ACCIDENT 06-0838648 70815 922556G
                Premium was applied as follows during the Plan/Policy Year -
                922556G ADD-BAS $7,507.49 159
                922556G LIFE-BTRM $30,562.30 171
                922556G LTD-ABIL $36,611.49 171
                Total $173,344.57
                """,
            ),
            (
                3,
                """
                The Hartford Group Benefits Division
                Annual Statement of Premiums and Producer Compensation
                HARTFORD LIFE AND ACCIDENT
                Producer and Address Org
                Code Policy
                Number Commissions
                Paid Fees Paid (1)Bonus
                Paid
                (2)Additional
                Compensation
                Paid
                NFP CORPORATE SERVICES
                NY LLC
                200 PARK AVE STE 3202
                NEW YORK, NY 10166
                3 922556-0GL $3,505.38 $0.00 $0.00 $0.00
                922556-GLT $4,008.18 $0.00 $0.00 $0.00
                Total $11,158.40 $0.00 $0.00 $0.00
                Producer and Address Org
                Code Policy
                Number Commissions
                Paid Fees Paid (1)Bonus
                Paid
                (2)Additional
                Compensation
                Paid
                NFP CORPORATE SERVICES
                NY LLC
                P.O. BOX 786677
                PHILADELPHIA, PA 19178
                3 922556-0GL $0.00 $0.00 $1,369.09 $0.00
                Total $0.00 $0.00 $4,324.75 $0.00
                (1)Bonus Paid represents contingent compensation.
                """,
            ),
        ]

        rows = extract_hartford_broker_rows(pages)
        fields = {field.field_name: field.value for field in extract_hartford_schedule_a_fields(pages)}

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].commission_total, "11,158.40")
        self.assertEqual(rows[1].fee_total, "4,324.75")
        self.assertEqual(rows[1].fee_rows[0].purpose, "BONUS PAID")
        self.assertEqual(fields["1a. Name of Insurance Company"], "HARTFORD LIFE AND ACCIDENT")
        self.assertEqual(fields["1b. Insurance Carrier EIN"], "06-0838648")
        self.assertEqual(fields["1c. NAIC Code"], "70815")
        self.assertEqual(fields["1d. Contract/Policy Number"], "922556G")
        self.assertEqual(fields["1e. Persons Covered (End of Policy Year)"], "171")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "173,344.57")
        self.assertEqual(fields["3b. Amount of Commissions"], "11,158.40")
        self.assertEqual(fields["3c. Amount of Fees"], "4,324.75")

    def test_image_only_hartford_uses_groundx_page_text_for_authoritative_layout_parser(self):
        identity_page = """
        The Hartford Group Benefits Division
        Annual Statement of Premiums and Producer Compensation
        Plan/Policy Year - 01/01/2025 to 12/31/2025
        Name of Insurance Carrier EIN NAIC Code Policy Number
        HARTFORD LIFE AND ACCIDENT 06-0838648 70815 922556G
        Premium was applied as follows during the Plan/Policy Year -
        922556G ADD-BAS $7,507.49 159
        922556G LIFE-BTRM $30,562.30 171
        Total $173,344.57
        """
        broker_page = """
        The Hartford Group Benefits Division
        Annual Statement of Premiums and Producer Compensation
        HARTFORD LIFE AND ACCIDENT
        Producer and Address Org
        Code Policy
        Number Commissions
        Paid Fees Paid (1)Bonus
        Paid
        (2)Additional
        Compensation
        Paid
        NFP CORPORATE SERVICES
        NY LLC
        P.O. BOX 786677
        PHILADELPHIA, PA 19178
        3 922556-0GL $0.00 $0.00 $1,369.09 $0.00
        Total $0.00 $0.00 $4,324.75 $0.00
        (1)Bonus Paid represents contingent compensation.
        """
        groundx = NormalizedExtractionResult(
            provider="GroundX X-Ray",
            fields=[],
            raw={
                "chunks": [
                    {"pageNumbers": [2], "suggestedText": identity_page},
                    {"pageNumbers": [3], "suggestedText": broker_page},
                ]
            },
        )
        settings = SimpleNamespace(
            schedule_a_canonical_validation_enabled=False,
            schedule_a_canonical_validation_shadow_enabled=True,
        )
        service = ExtractionService()

        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.extract_document_text_pages", return_value=[]),
            patch("app.services.extractor.extract_pdf_layout_text_pages", return_value=[]),
            patch.object(service, "_extract_schedule_a_unresolved", new=AsyncMock(return_value=groundx)),
        ):
            result = asyncio.run(service.extract_schedule_a(b"image-only-pdf", "4. HARTFORD.pdf"))

        fields = {field.field_name: field.value for field in result.fields}
        self.assertEqual(fields["1a. Name of Insurance Company"], "HARTFORD LIFE AND ACCIDENT")
        self.assertEqual(fields["1d. Contract/Policy Number"], "922556G")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "173,344.57")
        self.assertEqual(len(result.schedule_a_broker_rows), 1)
        self.assertEqual(result.schedule_a_broker_rows[0].fee_total, "4,324.75")
        self.assertEqual(result.schedule_a_broker_rows[0].source_page, 3)

    def test_allone_eap_compact_doc_extracts_identity_period_lives_and_paid_total(self):
        pages = [
            (
                1,
                """
                AllOne Health EAP
                Form 5500 Schedule A Information for: CareQuest
                Service Period = 1/1/25-12/31/25
                Headcount = 178
                Monthly rate = $2.78 per employee per month
                Total administrative fees paid = $5,938.08
                """,
            )
        ]

        fields = {field.field_name: field.value for field in extract_allone_eap_schedule_a_fields(pages)}

        self.assertEqual(fields["1a. Name of Insurance Company"], "ALLONE HEALTH")
        self.assertEqual(fields["1e. Persons Covered (End of Policy Year)"], "178")
        self.assertEqual(fields["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(fields["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "5,938.08")

    def test_metlife_bay_bridge_multirecord_report_preserves_both_records_and_brokers(self):
        def page(carrier, ein, naic, product, premium, employees, dependents, marsh, nfp):
            return f"""
            Insurance Data for Schedule A – Form 5500
            Name of Carrier: {carrier}
            Carrier EIN: {ein}
            Carrier NAIC Code: {naic}
            Group/account Name: Carequest Institute for Oral Health, Inc.
            Product Type: {product}
            Year: 1/01/2025 – 12/31/2025
            Total Premium at Year End: ${premium}
            Total # of Employees: {employees}
            Total # of Dependent: {dependents}
            Insurance Commission Information for Schedule A – Form 5500
            Producer Name and Address: Commissions paid to producers:
            Marsh & McLennan Agency, LLC. ${marsh}
            101 Huntington Ave. Ste. 401
            Boston, MA 21997
            NFP Corporate Services NY, LLC ${nfp}
            200 Park Ave. Rm 3202
            New York, NY 10166
            """

        pages = [
            (1, page("MetLife Legal Plan", "341650967", "", "Legal Plan", "3,149.97", "13", "18", "167.76", "109.28")),
            (2, page("MetLife Insurance Company", "135581829", "65978", "Accident", "6,528.09", "18", "15", "1,179.87", "125.95")),
        ]

        summaries = extract_metlife_bay_bridge_schedule_a_summaries(pages)
        brokers = extract_metlife_bay_bridge_broker_rows(pages)
        fields = {field.field_name: field for field in extract_metlife_bay_bridge_schedule_a_fields(pages)}

        self.assertEqual(len(summaries), 2)
        self.assertEqual(summaries[0].carrier_name, "MetLife Legal Plan")
        self.assertEqual(summaries[0].ein, "34-1650967")
        self.assertEqual(summaries[0].values[0].value, "31")
        self.assertEqual(summaries[1].naic_code, "65978")
        self.assertEqual(summaries[1].values[1].value, "6,528.09")
        self.assertEqual(len(brokers), 4)
        self.assertEqual(brokers[0].commission_rows[0].coverage, "Legal Plan")
        self.assertEqual(brokers[2].commission_rows[0].coverage, "Accident")
        self.assertEqual(fields["1a. Name of Insurance Company"].decision, "REVIEW_REQUIRED")
        self.assertEqual(
            fields["1a. Name of Insurance Company"].candidate_values,
            ["MetLife Legal Plan", "MetLife Insurance Company"],
        )

    def test_nonexperience_section_drops_false_experience_rated_commission(self):
        result = NormalizedExtractionResult(
            provider="test",
            fields=[
                NormalizedExtractionField(field_name="3b. Amount of Commissions", value="100.00", confidence=0.99),
                NormalizedExtractionField(field_name="9c(1)(A). Commissions", value="100.00", confidence=0.92),
                NormalizedExtractionField(
                    field_name="10a. Total premiums or subscription charges paid to carrier",
                    value="5,000.00",
                    confidence=0.99,
                ),
            ],
        )

        cleaned = remove_inapplicable_experience_rated_fields(
            result,
            [(1, "Part III Welfare Benefit Contract Information\n9. Non experience-rated contracts")],
        )

        self.assertEqual(
            [field.field_name for field in cleaned.fields],
            ["3b. Amount of Commissions", "10a. Total premiums or subscription charges paid to carrier"],
        )

    def test_unitedhealthcare_part_i_extracts_structured_broker_for_existing_row_match(self):
        pages = [
            (
                2,
                """
                Schedule A (Form 5500) Parts I and III
                Insurance Information Certified by Carrier
                (a) Name of Insurance carrier: UnitedHealthcare Insurance Company
                2. Insurance fees and commissions paid to agents, brokers, and other persons
                (a) Name and address of the agents, brokers or other persons to whom commissions or fees were paid:
                NFP CORPORATE SERVICES (MA) LLC
                141 LONGWATER DR STE 101
                NORWELL MA 02061-1620
                (b) Amount of commissions paid: $100,901.12
                (c) Fees paid / Amount: $0.00
                (d) Fees paid/Purpose: N/A
                (e) Organizational Code: 3
                """,
            )
        ]

        rows = extract_unitedhealthcare_broker_rows(pages)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "NFP CORPORATE SERVICES (MA) LLC")
        self.assertEqual(rows[0].address_line_1, "141 LONGWATER DR STE 101")
        self.assertEqual(rows[0].city, "NORWELL")
        self.assertEqual(rows[0].state, "MA")
        self.assertEqual(rows[0].zip_code, "02061-1620")
        self.assertEqual(rows[0].commission_total, "100,901.12")
        self.assertEqual(rows[0].fee_total, "0.00")
        self.assertEqual(rows[0].organization_code, "3")

    def test_vsp_floor_address_does_not_shift_boston_into_street_or_city(self):
        pages = [
            (
                1,
                """
                VSP vision care Schedule A Form (5500) Insurance Information
                Insurance Carrier: Vision Service Plan
                Commissions/Fees Paid for Policy Agent or Broker or Contract Year
                RSC Insurance Brokerage, Inc. $1,275.91160 Federal St 4th FloorBOSTON MA 02110
                """,
            )
        ]

        rows = extract_vsp_broker_rows(pages)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].address_line_1, "160 Federal St 4th Floor")
        self.assertEqual(rows[0].city, "BOSTON")

    def test_aetna_attached_listing_uses_plan_sponsor_state_and_compact_broker_row(self):
        pages = [
            (
                1,
                """
                VITCO FOODS
                715 E CALIFORNIA ST
                ONTARIO CA       91761
                Aetna Health, Inc.
                """,
            ),
            (
                3,
                """
                AETNA LIFE INSURANCE COMPANY
                2. Insurance Fees and commissions paid to agents and brokers:
                Contract or (a) Name and address of the agents or brokers (b) Amount of
                Identification to whom commissions or fees were paid. commissions paid
                0252023HNO JASON ANDREW PRATTES 8182 NOELLE DR $43,139.39
                HUNTINGTON BEACH, CA 92646
                """,
            ),
            (
                4,
                """
                State NAIC Code Service Area EIN
                AZ 95109 Aetna Health Inc. (a Pennsylvania Corporation) 23-2169745
                CA 00000 Aetna Health of California Inc. (a California Corporation) 95-3402799
                """,
            ),
        ]

        attached = {field.field_name: field.value for field in extract_aetna_attached_listing_fields(pages)}
        self.assertEqual(attached["1a. Name of Insurance Company"], "Aetna Health of California Inc.")
        self.assertEqual(attached["1b. Insurance Carrier EIN"], "95-3402799")
        self.assertEqual(attached["1c. NAIC Code"], "00000")

        brokers = extract_position_aware_schedule_a_broker_rows(pages)
        self.assertEqual(len(brokers), 1)
        self.assertEqual(brokers[0].name, "JASON ANDREW PRATTES")
        self.assertEqual(brokers[0].commission_total, "43,139.39")
        self.assertEqual(brokers[0].fee_total, "0")
        self.assertEqual(brokers[0].state, "CA")

    def test_aetna_support_statement_uses_the_matching_naic_appendix_row(self):
        pages = [
            (
                3,
                """
                AETNA LIFE INSURANCE COMPANY
                (a) Name of Insurance Carrier: (d) Contract Number
                Aetna Life Insurance Co. or Identification:
                (b) EIN: 06-6033492 0252023 of policy or contract year
                """,
            ),
            (
                4,
                """
                NAIC Code Service Area
                95094 Aetna Health Inc. (a Georgia Corporation)
                60054 Aetna Life Insurance Company
                """,
            ),
        ]
        by_name = {
            field.field_name: field.value
            for field in extract_aetna_schedule_a_support_statement_fields(pages)
        }

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Aetna Life Insurance Co.")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "06-6033492")
        self.assertEqual(by_name["1c. NAIC Code"], "60054")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "0252023")

    def test_litera_aetna_medical_uses_part_i_values_and_ignores_schedule_c_compensation(self):
        pages = [
            (
                3,
                """
                AETNA LIFE INSURANCE COMPANY AND AFFILIATES
                For Fiscal Plan Year beginning 01/01/2025 and ending 12/31/2025
                (a) Name of Insurance Carrier: (d) Contract Number (e) Approximate Number of
                Aetna Life Insurance Co. or Identification: persons covered at the end Policy or contract Year
                (b) EIN: 06-6033492 0186483-Medical of policy or contract year: (f) From: (g) To:
                (c) NAIC Code: See Attached Listing 685 01/01/2025 12/31/2025
                2. Insurance Fees and commissions paid to agents and brokers:
                Contract or (a) Name and address of the agents or brokers (b) Amount of (c) & (d) Fees Paid
                Identification to whom commissions or fees were paid. commissions paid Amount Purpose
                Reported fees and commissions may be attributed to multiple Aetna companies.
                Health (other than dental or vision)
                9. Non experience rated contracts:
                (a) Total premiums or subscription Charges paid to Carriel $5,693,303.00
                """,
            ),
            (
                5,
                """
                completing Schedule C of Form 5500.
                Broker Information: MERCER HEALTH & BENEFITS LLC
                4565 PAYSPHERE CIRCLE CHICAGO, IL 60674
                $240,148.74
                """,
            ),
        ]

        values = {
            field.field_name: field.value
            for field in extract_litera_aetna_schedule_a_fields(pages)
        }
        rows = extract_litera_aetna_schedule_a_broker_rows(pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "Aetna Life Insurance Co.")
        self.assertEqual(values["1b. Insurance Carrier EIN"], "06-6033492")
        self.assertEqual(values["1c. NAIC Code"], "60054")
        self.assertEqual(values["1d. Contract/Policy Number"], "0186483-Medical")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "685")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "5,693,303.00")
        self.assertEqual(rows, [])

    def test_litera_aetna_indemnity_extracts_primary_schedule_a_broker_row(self):
        pages = [
            (
                3,
                """
                AETNA LIFE INSURANCE COMPANY
                For Fiscal Plan Year beginning 01/01/2025 and ending 12/31/2025
                (a) Name of Insurance Carrier: (d) Contract Number (e) Approximate Number of
                Aetna Life Insurance Co. or Identification: persons covered at the end Policy or contract Year
                (b) EIN: 06-6033492 803136 of policy or contract year: 211 (f) From: (g) To:
                (c) NAIC Code: See Attached 01/01/2025 12/31/2025
                2. Insurance Fees and commissions paid to agents and brokers:
                803136 MERCER HEALTH & BENEFITS ADMINISTRATION LLC 12421 MEREDITH DR $21,871.85
                URBANDALE IA 50398-900
                TOTAL $21,871.85
                Indemnity contract
                9. Non experience rated contracts:
                (a) Total premiums or subscription charges paid to carrier $45,873.97
                """,
            )
        ]

        values = {
            field.field_name: field.value
            for field in extract_litera_aetna_schedule_a_fields(pages)
        }
        rows = extract_litera_aetna_schedule_a_broker_rows(pages)

        self.assertEqual(values["1c. NAIC Code"], "60054")
        self.assertEqual(values["1d. Contract/Policy Number"], "803136")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "211")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "45,873.97")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "MERCER HEALTH & BENEFITS ADMINISTRATION LLC")
        self.assertEqual(rows[0].address_line_1, "12421 MEREDITH DR")
        self.assertEqual(rows[0].city, "URBANDALE")
        self.assertEqual(rows[0].state, "IA")
        self.assertEqual(rows[0].zip_code, "50398-0900")
        self.assertEqual(rows[0].commission_total, "21,871.85")
        self.assertEqual(rows[0].fee_total, "0")
        self.assertEqual(rows[0].purpose, "COMMISSIONS")

    def test_litera_aetna_flattened_native_text_keeps_lives_premium_and_commission_columns(self):
        pages = [
            (
                3,
                "AETNA LIFE INSURANCE COMPANY AND AFFILIATES "
                "For Fiscal Plan Year beginning 01/01/2025 and ending 12/31/2025 "
                "1. Coverage: Traditional Prospective "
                "(a) Name of Insurance Carrier: Aetna Life Insurance Co."
                "(d) Contract Number or Identification:"
                "(e) Approximate Number of persons covered at the end of policy or contract year: "
                "Policy or contract Year(b) EIN: 06-6033492 803136 (f) From:(g)To:"
                "(c) NAIC Code: See Attached211 01/01/202512/31/2025 "
                "2. Insurance Fees and commissions paid to agents and brokers:"
                "803136MERCER HEALTH & BENEFITSADMINISTRATION LLC 12421 MEREDITH DR"
                "URBANDALE IA 50398-900$21,871.85 TOTAL $21,871.85 "
                "Part III Welfare Benefit Contract Information 9. Non experience rated contracts:"
                "(a) Total premiums or subscription charges paid to carrier $45,873.97",
            )
        ]

        fields = extract_litera_aetna_schedule_a_fields(pages)
        rows = extract_litera_aetna_schedule_a_broker_rows(pages)
        values = {field.field_name: field.value for field in fields}
        resolved = resolve_schedule_a_result(
            NormalizedExtractionResult(
                provider="Aetna filled Schedule A parser",
                fields=fields,
                schedule_a_broker_rows=rows,
                raw={},
            )
        )

        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "211")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "45,873.97")
        self.assertEqual(rows[0].commission_total, "21,871.85")
        self.assertEqual(rows[0].fee_total, "0")
        self.assertTrue(all(field.decision == "AUTOMATIC" for field in resolved.fields))
        self.assertEqual(resolved.schedule_a_broker_rows[0].decision, "AUTOMATIC")

    def test_litera_lincoln_voluntary_life_preserves_contract_suffix_and_merges_broker_payments(self):
        pages = [
            (
                2,
                """
                THE LINCOLN NATIONAL LIFE INSURANCE COMPANY
                SCHEDULE A REPORTING INFORMATION
                Part I - Information Concerning Insurance Contract Coverage, Fees, and Commissions
                (a) Name of insurance carrier: The Lincoln National Life Insurance Company
                (bo) EIN: 35-0472300
                (c) NAIC code: 65676
                (d) Contract or identification number: 000400001000 23109
                (Part III, #8) (e) (f) (g)
                Vol Child Life 36 01/01/2025 12/31/2025
                Vol Spouse AD&D 42 01/01/2025 12/31/2025
                Vol Spouse Life 42 01/01/2025 12/31/2025
                Voluntary AD&D 88 01/01/2025 12/31/2025
                Voluntary Life 88 01/01/2025 12/31/2025
                2. Insurance fee and commission information.
                (a) Total amount of commissions paid (b) Total amount of fees paid
                $5,938.27 $549.77
                3. Insurance fees and commissions paid to agents, brokers, and other persons:
                MERCER HEALTH & BENEFITS LLC $5,938.27 3
                4565 PAYSPHERE CIR
                CHICAGO, IL 60674
                Totals: $5,938.27 $0.00
                MERCER HEALTH & BENEFITS LLC $549.77 Broker Bonus 3
                4565 PAYSPHERE CIRCLE
                CHICAGO, IL 60674
                Totals: $0.00 $549.77
                Part III - Welfare Benefit Contract Information
                10. Nonexperience-rated contracts:
                (a) Total premiums or subscription charges paid to carrier $39,588.17
                """,
            )
        ]

        values = {
            field.field_name: field.value
            for field in extract_litera_lincoln_schedule_a_fields(pages)
        }
        rows = extract_litera_lincoln_schedule_a_broker_rows(pages)

        self.assertEqual(values["1d. Contract/Policy Number"], "00040000100023109")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "88")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "39,588.17")
        self.assertEqual(values["3b. Amount of Commissions"], "5,938.27")
        self.assertEqual(values["3c. Amount of Fees"], "549.77")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].commission_total, "5,938.27")
        self.assertEqual(rows[0].fee_total, "549.77")
        self.assertEqual(rows[0].fee_rows[0].purpose, "Broker Bonus")
        resolved = resolve_schedule_a_result(
            NormalizedExtractionResult(
                provider="Lincoln Schedule A reporting parser",
                fields=extract_litera_lincoln_schedule_a_fields(pages),
                schedule_a_broker_rows=extract_litera_lincoln_schedule_a_broker_rows(pages),
                raw={},
            )
        )
        self.assertEqual(
            next(field for field in resolved.fields if field.field_name.startswith("10a.")).decision,
            "AUTOMATIC",
        )
        self.assertEqual(resolved.schedule_a_broker_rows[0].decision, "AUTOMATIC")
        self.assertFalse(
            any(
                item.validator == "broker_total_reconciliation" and item.status == "ERROR"
                for item in resolved.schedule_a_broker_rows[0].validation_results
            )
        )

    def test_litera_lincoln_image_ocr_accepts_joined_headings_and_preserves_full_policy(self):
        pages = [
            (
                2,
                """
                THELINCOLNNATIONALLIFEINSURANCECOMPANY
                SCHEDULEAREPORTINGINFORMATION
                (a) Name of insurance carrier: The Lincoln National Life Insurance Company
                EIN:35-0472300
                NAIC code: 65676
                (d) Contract or identification number: 000010233868 00000
                (Part III, #8) (e) (f) (g)
                STD 422 01/01/2025 12/31/2025
                2. Insurance fee and commission information.
                10. Nonexperience-rated contracts:
                (a) Total premiums or subscription charges paid to carrier $126,256.23
                """,
            )
        ]

        values = {
            field.field_name: field.value
            for field in extract_litera_lincoln_schedule_a_fields(pages)
        }

        self.assertEqual(values["1b. Insurance Carrier EIN"], "35-0472300")
        self.assertEqual(values["1c. NAIC Code"], "65676")
        self.assertEqual(values["1d. Contract/Policy Number"], "00001023386800000")

    def test_local_page_pipeline_includes_litera_lincoln_authoritative_fields(self):
        pages = [
            (
                2,
                """
                THELINCOLNNATIONALLIFEINSURANCECOMPANY
                SCHEDULEAREPORTINGINFORMATION
                (a) Name of insurance carrier: The Lincoln National Life Insurance Company
                EIN:35-0472300
                NAIC code: 65676
                (d) Contract or identification number: 000010233867 00000
                (Part III, #8) (e) (f) (g)
                LTD 422 01/01/2025 12/31/2025
                2. Insurance fee and commission information.
                10. Nonexperience-rated contracts:
                (a) Total premiums or subscription charges paid to carrier $126,256.23
                """,
            )
        ]

        values = {field.field_name: field.value for field in _extract_fields_from_pages(pages)}

        self.assertEqual(values["1d. Contract/Policy Number"], "00001023386700000")

    def test_litera_lincoln_recovers_dotted_premium_from_xray_glyph_noise(self):
        pages = [
            (
                2,
                """
                THE LINCOLN NATIONAL LIFE INSURANCE COMPANY
                SCHEDULE A REPORTING INFORMATION
                (a) Name of insurance carrier: The Lincoln National Life Insurance Company
                (b) EIN: 35-0472300
                (c) NAIC code: 65676
                (d) Contract or identification number: 000010233867 00000
                (Part III, #8) (e) (f) (g)
                LTD 422 01/01/2025 12/31/2025
                2. Insurance fee and commission information.
                Part III - Welfare Benefit Contract Information
                10. Nonexperience-rated contracts:
                (a) Total premiums or subscription charges paid to carrier... S126, 256523
                """,
            )
        ]

        values = {
            field.field_name: field.value
            for field in extract_litera_lincoln_schedule_a_fields(pages)
        }

        self.assertEqual(
            values["10a. Total premiums or subscription charges paid to carrier"],
            "126,256.23",
        )

    def test_principal_short_form_extracts_part_i_and_broker_values(self):
        pages = [
            (1, """
                C ontract # 1205934
                D ata Period September 1, 2025 to December 31, 2025
                P rincipal Life Insurance Company
                (b) EIN 42-0127290 (c) NAIC Code 61271
                3 62 E mployees
                Section 10: Non-Experience Rated Contracts (a) Total Premiums Paid to Carrier 1 4,655
                Total (from below) (a) Commissions Paid 1 ,897
                JASON ANDREW PRATTES
                620 NEWPORT CENTER DR STE 1100
                NEWPORT BEACH CA 92660-8011 1 ,897 3 - Ins Agent or Broker
            """),
        ]

        fields = extract_principal_short_form_schedule_a_fields(pages)
        values = {field.field_name: field.value for field in fields}
        rows = extract_principal_short_form_broker_rows(pages)

        self.assertEqual(values["1d. Contract/Policy Number"], "1205934")
        self.assertEqual(values["1e. Persons Covered (End of Policy Year)"], "362")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "14,655")
        self.assertEqual(values["3b. Amount of Commissions"], "1,897")
        contract = next(field for field in fields if field.field_name == "1d. Contract/Policy Number")
        self.assertIn("1205934", contract.source_text or "")
        self.assertEqual(rows[0].name, "JASON ANDREW PRATTES")
        self.assertEqual(rows[0].commission_total, "1,897")
        self.assertEqual(rows[0].organization_code, "3")

    def test_principal_short_form_restores_labelled_values_after_semantic_enrichment(self):
        pages = [(1, """
            C ontract # 1205934
            D ata Period September 1, 2025 to December 31, 2025
            P rincipal Life Insurance Company
            (b) EIN 42-0127290 (c) NAIC Code 61271
        """)]
        result = NormalizedExtractionResult(
            provider="test",
            fields=[
                NormalizedExtractionField(
                    field_name="1a. Name of Insurance Company",
                    value="Principal Life Insurance Company (b)",
                    confidence=0.99,
                ),
                NormalizedExtractionField(
                    field_name="1d. Contract/Policy Number",
                    value="1205934",
                    confidence=0.99,
                ),
            ],
        )

        restored = restore_principal_short_form_fields(result, pages)
        values = {field.field_name: field.value for field in restored.fields}

        self.assertEqual(values["1a. Name of Insurance Company"], "Principal Life Insurance Company")
        self.assertEqual(values["1d. Contract/Policy Number"], "1205934")
        self.assertEqual(values["1f. Policy Year Beginning Date"], "09/01/2025")
        self.assertEqual(values["1g. Policy Year Ending Date"], "12/31/2025")

    def test_standard_short_form_extracts_part_i_and_broker_values(self):
        pages = [
            (1, """
                JASON PRATTES
                1947 PORT LAURENT PL
                NEWPORT BEACH, CA 92660
                $3,638.32 $0.00 $0.00 $0.00 3
                TOTAL COMMISSIONS PAID $3,638.32
                TOTAL CONTINGENT COMP PAID $0.00
                Anthem Life Insurance Company
                VITCO DISTRIBUTORS
                1/1/2025
                12/31/2025
                0
                35-0980405
                000-61069
                LIFE INSURANCE
                PLAN INFORMATION REPORT FOR THE PERIOD OF
                282802
                SHORT FORM INFORMATION
            """),
            (2, """
                (a) TOTAL PREMIUM PAID TO CARRIER:
                Anthem Life Insurance Company HEREBY CERTIFIES THAT THIS INFORMATION IS COMPLETE
                282802
                LIFE INSURANCE
                $33,126.32
                SHORT FORM INFORMATION
            """),
        ]

        values = {field.field_name: field.value for field in extract_standard_short_form_schedule_a_fields(pages)}
        rows = extract_standard_short_form_broker_rows(pages)

        self.assertEqual(values["1a. Name of Insurance Company"], "Anthem Life Insurance Company")
        self.assertEqual(values["1c. NAIC Code"], "61069")
        self.assertEqual(values["1d. Contract/Policy Number"], "282802")
        self.assertEqual(values["10a. Total premiums or subscription charges paid to carrier"], "33,126.32")
        self.assertEqual(rows[0].name, "JASON PRATTES")
        self.assertEqual(rows[0].commission_total, "3,638.32")
        self.assertEqual(rows[0].fee_total, "0.00")

    def test_money_value_normalizes_leading_decimal_zero(self):
        self.assertEqual(money_value("$.00"), "0.00")

    def test_labeled_naic_strips_only_leading_zero_padding_to_five_digits(self):
        fields = extract_schedule_a_fields_from_rule_labels(
            "NAIC Code: 00053295",
            page=1,
            rules=DEFAULT_FIELD_RULES,
        )

        by_name = {field.field_name: field.value for field in fields}
        self.assertEqual(by_name["1c. NAIC Code"], "53295")

    def test_labeled_naic_does_not_hide_a_nonzero_six_digit_source_error(self):
        fields = extract_schedule_a_fields_from_rule_labels(
            "NAIC Code: 624190",
            page=1,
            rules=DEFAULT_FIELD_RULES,
        )

        by_name = {field.field_name: field.value for field in fields}
        self.assertEqual(by_name["1c. NAIC Code"], "624190")

    def test_position_aware_naic_uses_full_zero_padded_source_value(self):
        fields = extract_position_aware_schedule_a_fields(
            [(1, "c. NAIC Code: 00053295")]
        )

        by_name = {field.field_name: field.value for field in fields}
        self.assertEqual(by_name["1c. NAIC Code"], "53295")

    def test_new_york_life_annual_policy_report_keeps_full_contract_and_totals(self):
        pages = [
            (
                2,
                """
                Anniversary
                Annual Policy Information Report
                Name of Insurance Carrier
                New Y ork Life Group Insurance Company of New Y ork
                EIN 23-1503749
                NAIC Code 65498
                Contract/Policy Number OK 0968358
                Contract/Policy Y ear From: 01/01/2025
                Contract/Policy Y ear T o: 12/31/2025
                T otal premiums paid to Insurance Company during the policy year: $ 7,462.96
                """,
            )
        ]

        by_name = {field.field_name: field.value for field in extract_nyl_annual_policy_fields(pages)}

        self.assertEqual(by_name["1a. Name of Insurance Company"], "New York Life Group Insurance Company of New York")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "23-1503749")
        self.assertEqual(by_name["1c. NAIC Code"], "65498")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "OK 0968358")
        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "7,462.96")

    def test_hmsa_support_packet_extracts_header_covered_count_and_total_premium(self):
        pages = [
            (
                1,
                """
                Attached is the HMSA Health Plan information that may be used to complete Schedule A.
                For filing purposes, please use EIN number 99-0040115 and NAIC code 49948.
                """,
            ),
            (
                2,
                """
                HAWAII MEDICAL SERVICE ASSOCIATION
                ERISA FORM 5500 AND SCHEDULE A INFORMATION
                FOR THE PLAN YEAR: January 2025 - December 2025
                Acct Code Group # Sub Group Group Name Subs Subs & Deps Paid Premium
                C000 012763 001 JTB USA INC 19 23 $210,730.54
                012763 002 JTB USA INC COBRA 0 0 $756.56
                """,
            ),
        ]

        by_name = {field.field_name: field.value for field in extract_hmsa_schedule_a_fields(pages)}

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Hawaii Medical Service Association (HMSA)")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "99-0040115")
        self.assertEqual(by_name["1c. NAIC Code"], "49948")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "012763")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "23")
        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "211,487.10")

    def test_schedule_a_parser_extracts_and_merges_columnar_broker_disclosure_rows(self):
        pages = [
            (
                2,
                """
                5. INSURANCE FEES AND COMMISSION INFORMATION:
                NAME AND ADDRESS OF EACH SOLICITING AGENT OR BROKER RECEIVING COMPENSATION:
                SALES COMMISSION PAID FEES PAID ADDITIONAL COMPENSATION PAID

                ALLIANT INSURANCE
                SERVICES, INC.
                $ 0.00 $ 0.00 $ 0.00
                32 OLD SLIP
                NEW YORK, NY 10005

                NFP CORPORATE SERVICES
                (NY) LLC
                $ 1,689.77 $ 0.00 $ 0.00
                PO BOX 9101
                PLAINVIEW, NY 11803

                USI INSURANCE SERVICES LLC $ 45.45 $ 0.00 $ 0.00
                3RD FLOOR
                600 THIRD AVENUE
                NEW YORK, NY 10016

                USI INSURANCE SERVICES LLC $ 0.00 $ 0.00 $ 52.98
                3RD FLOOR
                600 THIRD AVENUE
                NEW YORK, NY 10016

                MANAGEMENT COMPENSATION
                GROUP/NFP
                $ 0.00 $ 0.00 $ 145.16
                STE 200
                3445 PEACHTREE RD NE
                """,
            ),
            (
                3,
                """
                ATLANTA, GA 30326

                NFP CORPORATE SERVICES
                (NY) LLC
                $ 1,035.96 $ 0.00 $ 0.00
                PO BOX 9101
                PLAINVIEW, NY 11803

                NFP INSURANCE SERVICES,
                INC
                $ 0.00 $ 0.00 $ 510.90
                1250 CAPITAL OF TEXAS HWY
                BLDG 2 STE 125
                AUSTIN, TX 78746

                6. COVERAGE/BENEFITS PROVIDED: DISABILITY
                """,
            ),
        ]

        rows = extract_columnar_broker_compensation_rows(pages)
        by_name = {row.name: row for row in rows}

        self.assertEqual(len(rows), 4)
        self.assertNotIn("ALLIANT INSURANCE SERVICES, INC.", by_name)
        self.assertEqual(by_name["NFP CORPORATE SERVICES (NY) LLC"].commission_total, "2,725.73")
        self.assertEqual(by_name["NFP CORPORATE SERVICES (NY) LLC"].fee_total, "0")
        self.assertEqual(by_name["NFP CORPORATE SERVICES (NY) LLC"].zip_code, "11803")
        self.assertEqual(by_name["USI INSURANCE SERVICES LLC"].commission_total, "45.45")
        self.assertEqual(by_name["USI INSURANCE SERVICES LLC"].fee_total, "52.98")
        self.assertEqual(by_name["USI INSURANCE SERVICES LLC"].address_line_1, "600 THIRD AVENUE")
        self.assertEqual(by_name["USI INSURANCE SERVICES LLC"].address_line_2, "3RD FLOOR")
        self.assertEqual(by_name["MANAGEMENT COMPENSATION GROUP/NFP"].city, "ATLANTA")
        self.assertEqual(by_name["MANAGEMENT COMPENSATION GROUP/NFP"].fee_total, "145.16")
        self.assertEqual(by_name["NFP INSURANCE SERVICES INC"].fee_total, "510.90")

        fields = {field.field_name: field.value for field in schedule_a_broker_compensation_fields(rows)}
        self.assertEqual(fields["3b. Amount of Commissions"], "2,771.18")
        self.assertEqual(fields["3c. Amount of Fees"], "709.04")

    def test_columnar_broker_disclosure_fails_closed_when_a_paid_row_has_no_name(self):
        rows = extract_columnar_broker_compensation_rows(
            [
                (
                    1,
                    """
                    INSURANCE FEES AND COMMISSION INFORMATION:
                    SALES COMMISSION PAID FEES PAID ADDITIONAL COMPENSATION PAID

                    $ 100.00 $ 0.00 $ 0.00
                    1 MAIN STREET
                    BOSTON, MA 02110

                    6. COVERAGE/BENEFITS PROVIDED: LIFE
                    """,
                )
            ]
        )

        self.assertEqual(rows, [])

    def test_columnar_broker_disclosure_handles_comma_name_and_cross_page_address(self):
        rows = extract_columnar_broker_compensation_rows(
            [
                (
                    2,
                    """
                    5. INSURANCE FEES AND COMMISSION INFORMATION:
                    SALES COMMISSION PAID FEES PAID ADDITIONAL COMPENSATION PAID

                    NFP INSURANCE SERVICES, $ 0.00 $ 0.00 $ 1,867.72
                    """,
                ),
                (
                    3,
                    """
                    INC
                    1250 CAPITAL OF TEXAS HWY
                    BLDG 2 STE 125
                    AUSTIN, TX 78746

                    6. COVERAGE/BENEFITS PROVIDED: DISABILITY
                    """,
                ),
            ]
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "NFP INSURANCE SERVICES INC")
        self.assertEqual(rows[0].address_line_1, "1250 CAPITAL OF TEXAS HWY")
        self.assertEqual(rows[0].address_line_2, "BLDG 2 STE 125")
        self.assertEqual(rows[0].fee_total, "1,867.72")

    def test_verified_broker_table_replaces_incorrect_ai_broker_values(self):
        ai_fields = [
            NormalizedExtractionField(field_name="3a. Name of Agent/Broker/Person", value="March", confidence=0.99),
            NormalizedExtractionField(field_name="3b. Amount of Commissions", value="31", confidence=0.99),
        ]
        broker_fields = schedule_a_broker_compensation_fields(
            [
                ScheduleABrokerRow(
                    name="NFP CORPORATE SERVICES (NY) LLC",
                    commission_total="2,725.73",
                    fee_total="0",
                    commission_rows=[ScheduleABrokerMoneyRow(amount="2,725.73", purpose="Sales Commission")],
                    source_page=2,
                )
            ]
        )

        merged = {field.field_name: field.value for field in merge_schedule_a_fields(ai_fields, broker_fields)}

        self.assertEqual(merged["3a. Name of Agent/Broker/Person"], "NFP CORPORATE SERVICES (NY) LLC")
        self.assertEqual(merged["3b. Amount of Commissions"], "2,725.73")

    def test_numeric_schedule_a_field_drops_due_label_fragment(self):
        merged = merge_schedule_a_fields(
            [
                NormalizedExtractionField(
                    field_name="9c(2). Dividends or retroactive rate refunds due",
                    value="DUE",
                    confidence=0.99,
                )
            ],
            [],
        )

        self.assertEqual(merged, [])

    def test_numeric_broker_name_is_rejected_but_fee_is_preserved(self):
        selected = select_best_schedule_a_fields(
            [
                NormalizedExtractionField(
                    field_name="3a. Name of Agent/Broker/Person",
                    value="384.40",
                    confidence=0.99,
                ),
                NormalizedExtractionField(
                    field_name="9c(1)(B). Administrative service or other fees",
                    value="384.40",
                    confidence=0.99,
                ),
            ]
        )

        values = {field.field_name: field.value for field in selected}
        self.assertNotIn("3a. Name of Agent/Broker/Person", values)
        self.assertEqual(values["9c(1)(B). Administrative service or other fees"], "384.40")

        mapped = map_extraction_to_rules(
            "filing",
            [
                NormalizedExtractionField(
                    field_name="3a. Name of Agent/Broker/Person",
                    value="384.40",
                    confidence=0.99,
                )
            ],
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )
        broker = next(
            field
            for field in mapped["fields"]
            if field.mapped_rule_key == "schedule_a_part_i_3a_name_of_agent_broker_person"
        )
        self.assertEqual(broker.proposed_value, "")
        self.assertEqual(broker.status.value, "MISSING")

    def test_authoritative_local_broker_table_replaces_provider_rows(self):
        provider = NormalizedExtractionResult(
            provider="GroundX",
            fields=[],
            schedule_a_broker_rows=[
                ScheduleABrokerRow(
                    name="NFP CORPORATE SERVICES NY LLC",
                    address_line_1="$860.74",
                    commission_total="860.74",
                    fee_total="0",
                    confidence=0.99,
                )
            ],
        )
        local_row = ScheduleABrokerRow(
            name="NFP CORPORATE SERVICES NY LLC",
            address_line_1="PO BOX 9101",
            city="PLAINVIEW",
            state="NY",
            zip_code="11803",
            commission_total="2,762.56",
            fee_total="860.74",
            confidence=0.96,
        )
        local = NormalizedExtractionResult(
            provider="Local PDF parser",
            fields=[],
            raw={"authoritative_broker_table": True},
            schedule_a_broker_rows=[local_row],
        )

        result = supplement_schedule_a_result_with_local(provider, local)

        self.assertEqual(result.schedule_a_broker_rows, [local_row])

    def test_authoritative_zero_compensation_table_removes_provider_broker_name(self):
        provider = NormalizedExtractionResult(
            provider="GroundX",
            fields=[
                NormalizedExtractionField(
                    field_name="3a. Name of Agent/Broker/Person",
                    value="Arthur J. Gallagher Risk Management Services LLC",
                    confidence=0.99,
                )
            ],
        )
        local = NormalizedExtractionResult(
            provider="Local PDF parser",
            fields=[],
            raw={"authoritative_broker_table": True},
            schedule_a_broker_rows=[
                ScheduleABrokerRow(
                    name="Arthur J. Gallagher Risk Management Services LLC",
                    commission_total="0",
                    fee_total="0.00",
                )
            ],
        )

        result = supplement_schedule_a_result_with_local(provider, local)

        self.assertNotIn(
            "3a. Name of Agent/Broker/Person",
            {field.field_name for field in result.fields},
        )

    def test_cigna_summary_page_wins_over_state_appendices(self):
        pages = [
            (
                1,
                """
                Cigna Health and Life Insurance Company
                Schedule A Insurance Information
                Part I Information Concerning Insurance Contract Coverage, Fees and Commissions
                (Summary of All Insurance Contracts Included in Part III)
                1. Coverage Information (a) Name of Insurance Carrier: Cigna Health and Life Insurance Company and affiliates (\"Cigna\")
                (b) EIN
                59-1031071
                (c) NAIC Code
                67369
                (d) Contract or Identification Number
                3341244
                (e) Approx. no. of persons covered at end of policy or contract year
                475 Employees
                Policy/Contract Year
                (f) From (g) To
                01/01/2025 12/31/2025
                2. Insurance fees and commissions information.
                (a) Total Amount of commissions paid $18,603 (b) Total Amount of fees paid $1,397
                3. Persons receiving commissions and fees.
                Non Experience - Rated
                NFP CORPORATE SERVICES (NY), LLC,
                PO BOX 786677, PHILADELPHIA, PA, $18,603 $1,397 General Agent Payments 3-Insurance Agent or Broker
                19178
                Part III Welfare Benefit Contract Information
                9. Experience-Rated Contracts This section not applicable for this Plan
                10. Nonexperience-rated contracts
                (a) Total premiums or subscriptions charges paid to carrier $375,747
                PART IV Provision of Information
                11. Did the insurance company fail to provide any information necessary to complete Schedule A? Yes No
                12. If the answer to line 11 is \"Yes\", specify the information not provided. Answer \"Not Applicable\"
                """,
            ),
            (
                2,
                """
                Appendix to 1a, b and c
                Cigna Health and Life Insurance Company
                06-1141174 95660 3341244 1 Employees 01/01/2025 12/31/2025
                """,
            ),
        ]

        fields = extract_cigna_schedule_a_fields(pages)
        by_name = {field.field_name: field.value for field in fields}
        rows = extract_cigna_schedule_a_broker_rows(pages)

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Cigna Health and Life Insurance Company and affiliates")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "59-1031071")
        self.assertEqual(by_name["1c. NAIC Code"], "67369")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "3341244")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "475")
        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(by_name["3b. Amount of Commissions"], "18,603")
        self.assertEqual(by_name["3c. Amount of Fees"], "1,397")
        self.assertEqual(by_name["3d. Purpose"], "General Agent Payments")
        self.assertEqual(by_name["3e. Organizational Code"], "3")
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "375,747")
        self.assertEqual(
            by_name["11. Did the insurance company fail to provide any information necessary to complete Schedule A?"],
            "No",
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "NFP CORPORATE SERVICES (NY), LLC")
        self.assertEqual(rows[0].address_line_1, "PO BOX 786677")
        self.assertEqual(rows[0].city, "PHILADELPHIA")
        self.assertEqual(rows[0].state, "PA")
        self.assertEqual(rows[0].zip_code, "19178")
        self.assertEqual(rows[0].commission_total, "18,603")
        self.assertEqual(rows[0].fee_total, "1,397")

        with (
            patch("app.services.extractor.extract_pdf_text_pages", return_value=pages),
            patch("app.services.extractor.extract_fields_from_document_text", return_value=[]),
        ):
            local_result = local_schedule_a_pdf_result(b"%PDF", "cigna.pdf")
        self.assertTrue(local_result.raw["authoritative_broker_table"])
        self.assertEqual(len(local_result.schedule_a_broker_rows), 1)
        self.assertEqual(local_result.schedule_a_broker_rows[0].commission_total, "18,603")

    def test_cigna_summary_identity_is_restored_after_final_candidate_selection(self):
        pages = [
            (
                2,
                """
                Cigna Health and Life Insurance Company
                Schedule A Insurance Information
                (Summary of All Insurance Contracts Included in Part III)
                2. Insurance fees and commissions information.
                Part III Welfare Benefit Contract Information
                (a) Name of Insurance Carrier: Cigna Health and Life Insurance Company and affiliates ("Cigna")
                (b) EIN 59-1031071 (c) NAIC Code 67369
                (d) Contract or Identification Number 3333085
                (e) Approx. no. of persons covered at end of policy or contract year 246 Employees
                Policy/Contract Year (f) From (g) To 04/01/2025 03/31/2026
                """,
            ),
            (
                9,
                """
                Schedule A Insurance Information - Appendix to 1a, b and c
                Name: Cigna Dental Health Plan of Arizona, Inc.
                EIN Code: 86-0807222
                NAIC Code: 47013
                """,
            ),
        ]
        selected = [
            NormalizedExtractionField(
                field_name="1a. Name of Insurance Company",
                value="Cigna Health and Life Insurance Company",
                confidence=0.99,
                page=2,
            ),
            NormalizedExtractionField(
                field_name="1b. Insurance Carrier EIN",
                value="59-2600475",
                confidence=0.99,
                page=4,
            ),
            NormalizedExtractionField(
                field_name="1c. NAIC Code",
                value="47013",
                confidence=0.99,
                page=9,
            ),
            NormalizedExtractionField(
                field_name="1d. Contract/Policy Number",
                value="3333085",
                confidence=0.99,
                page=2,
            ),
        ]

        restored = prefer_authoritative_cigna_summary_fields(selected, pages)
        values = {field.field_name: field.value for field in restored}

        self.assertEqual(values["1b. Insurance Carrier EIN"], "59-1031071")
        self.assertEqual(values["1c. NAIC Code"], "67369")

    def test_cigna_schedule_a_support_packet_uses_plan_detail_not_schedule_c_disclosures(self):
        pages = [
            (
                2,
                """
                INFORMATION FOR COMPLETING SCHEDULE A ON THE IRS FORM 5500
                For Plan Year Beginning:
                January 01, 2025
                and Ending:
                December 31, 2025
                Name of Plan:
                JTB Americas, Ltd.
                SCHEDULE A - INSURANCE INFORMATION:
                Name of Insurance Carrier:
                Cigna Health and Life Insurance Company
                EIN
                NAIC
                Contract or identification
                Policy or contract year:
                To
                59-1031071
                67369
                00656053
                12/31/2025
                1/1/2025
                From
                Approximate number of persons covered at end of policy or contract year:
                Total premiums* or subscription charges paid to carrier:
                $891,167.11
                """,
            ),
            (
                4,
                """
                Eligible Indirect Compensation
                Service Provider Information for Reporting on Form 5500 Schedule C Part 1, Line 3
                (a) Service provider name: Cigna
                Eligible Indirect Compensation Formula/Estimate: For calendar year 2025, $0.07 PMPY
                Sources of indirect compensation, excluding eligible indirect compensation, to be reported on Schedule C Part 1, Line 3
                """,
            ),
            (
                6,
                """
                PREMIUMS PLAN DETAIL
                COMMISSIONS PAID DETAIL
                BENEFIT ADVISOR FEE PAID DETAIL
                SERVICE AND / OR GENERAL AGENT FEE PAID DETAIL
                BENEFIT
                TOTAL COMM PAID
                BROKER ACCT#
                BROKER NAME
                MEDICAL
                $192,508.68
                113447
                ALLIANCE 360 INSURANCE SOLUTIONS
                TOTAL
                $192,508.68
                BENEFIT
                TOTAL FEES
                BROKER ACCT#
                BROKER NAME
                TOTAL
                BENEFIT
                TOTAL FEES
                BROKER ACCT#
                BROKER NAME
                MEDICAL
                $77,014.08
                302347
                CENTERSTONE INS & FIN SVC LLC
                TOTAL
                $77,014.08
                For Plan Year Beginning:
                January 01, 2025
                and Ending:
                December 31, 2025
                Name of Plan:
                JTB Americas, Ltd.
                Cigna
                Plan Detail Report
                Plan #:
                00656053
                """,
            ),
        ]

        fields = {field.field_name: field.value for field in extract_cigna_schedule_a_fields(pages)}
        rows = extract_cigna_schedule_a_broker_rows(pages)

        self.assertEqual(fields["1a. Name of Insurance Company"], "Cigna Health and Life Insurance Company")
        self.assertEqual(fields["1b. Insurance Carrier EIN"], "59-1031071")
        self.assertEqual(fields["1c. NAIC Code"], "67369")
        self.assertEqual(fields["1d. Contract/Policy Number"], "00656053")
        self.assertEqual(fields["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(fields["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(fields["3b. Amount of Commissions"], "192,508.68")
        self.assertEqual(fields["3c. Amount of Fees"], "77,014.08")
        self.assertEqual(fields["10a. Total premiums or subscription charges paid to carrier"], "891,167.11")
        self.assertEqual([row.name for row in rows], ["ALLIANCE 360 INSURANCE SOLUTIONS", "CENTERSTONE INS & FIN SVC LLC"])
        self.assertEqual(rows[0].commission_total, "192,508.68")
        self.assertEqual(rows[0].fee_total, "0")
        self.assertEqual(rows[1].commission_total, "0")
        self.assertEqual(rows[1].fee_total, "77,014.08")

    def test_equitable_schedule_a_worksheet_extracts_coverage_period_and_checkbox_no_values(self):
        health_rule = FieldRule(
            key="ftw_discovered_schedule_a_health_ind",
            label="Health Indicator",
            ftw_field="Health Indicator",
            xml_tag="HealthInd",
            mapping_mode=FieldRuleMappingMode.FTW_MAPPED,
            priority="MEDIUM",
            source="Schedule A",
            form_section="Schedule A - Discovered FTW fields",
            field_type="Dynamic",
            existing_behavior="Review Only",
            new_behavior="Keep FTW",
            aliases=["Health"],
        )
        vision_rule = health_rule.model_copy(
            update={
                "key": "ftw_discovered_schedule_a_vision_ind",
                "label": "Vision Indicator",
                "ftw_field": "Vision Indicator",
                "xml_tag": "VisionInd",
                "aliases": ["Vision"],
            }
        )
        text = """
        Schedule A (Form5500)Worksheet
        (D) Contract or ID Number 011335 Total (E) 279 Combined Numbers
        Approx. no. of Persons cov. At End of Policy Year Employees
        (E) Policy or Contract Year (F) From (F) 2023-10-01 (G) To (G) 2024-09-30
        Section 8: Benefit and Contract Type
        (A) [] Health (other than dental or vision) (C) [] Vision (D) [X] Life Ins.
        """

        fields = parse_schedule_a_text(text, rules=[health_rule, vision_rule])
        fields.extend(
            extract_explicit_benefit_indicator_fields(
                text,
                rules=[health_rule, vision_rule],
            )
        )
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "279")
        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "10/01/2023")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "09/30/2024")
        self.assertEqual(by_name["Health Indicator"], "No")
        self.assertEqual(by_name["Vision Indicator"], "No")

    def test_obvious_irs_template_placeholders_are_not_real_values(self):
        self.assertTrue(is_obvious_template_placeholder("ABCDEFGHI ABCDEFGHI ABCDEFGHI"))
        self.assertTrue(is_obvious_template_placeholder("123456789012345"))
        self.assertFalse(is_obvious_template_placeholder("Federal Insurance Company"))
        self.assertFalse(is_obvious_template_placeholder("0927447"))

    def test_unfilled_irs_schedule_a_template_is_detected(self):
        pages = [
            (
                1,
                """
                SCHEDULE A (Form 5500) Insurance Information
                ABCDEFGHI ABCDEFGHI ABCDEFGHI ABCDEFGHI ABCDEFGHI ABCDEFGHI
                ABCDEFGHI ABCDEFGHI ABCDEFGHI
                012345678 ABCDE 123456789012345 123456789012345
                Policy or contract year YYYY-MM-DD YYYY-MM-DD
                Fees paid 123456789012345
                """,
            )
        ]

        self.assertTrue(is_unfilled_schedule_a_template(pages))

    def test_completed_overlay_is_not_rejected_when_template_layer_remains(self):
        pages = [
            (
                1,
                """
                SCHEDULE A (Form 5500) Insurance Information
                ABCDEFGHI ABCDEFGHI ABCDEFGHI ABCDEFGHI ABCDEFGHI ABCDEFGHI
                ABCDEFGHI ABCDEFGHI ABCDEFGHI
                012345678 ABCDE 123456789012345 123456789012345
                Policy or contract year YYYY-MM-DD YYYY-MM-DD
                01/01/2025 12/31/2025 Carrier EIN 13-4029115
                """,
            )
        ]

        self.assertFalse(is_unfilled_schedule_a_template(pages))

    def test_unfilled_template_preflight_skips_groundx(self):
        service = ExtractionService()
        pages = [
            (
                1,
                "SCHEDULE A (Form 5500) Insurance Information "
                + "ABCDEFGHI " * 9
                + "012345678 123456789012345 " * 3
                + "YYYY-MM-DD YYYY-MM-DD",
            )
        ]

        with (
            patch("app.services.extractor.extract_document_text_pages", return_value=pages),
            patch.object(service, "_extract_schedule_a_unresolved", new=AsyncMock()) as groundx,
        ):
            result = asyncio.run(service.extract_schedule_a(b"pdf", "blank-schedule-a.pdf"))

        groundx.assert_not_awaited()
        self.assertEqual(result.fields, [])
        self.assertTrue(result.raw["manual_review_required"])
        self.assertEqual(result.classification_signals, ["UNFILLED_SCHEDULE_A_TEMPLATE"])

    def test_local_parser_uses_discovered_ftw_aliases_for_explicit_benefits(self):
        health_rule = FieldRule(
            key="ftw_discovered_schedule_a_health_ind",
            label="Health Indicator",
            ftw_field="Health Indicator",
            xml_tag="HealthInd",
            mapping_mode=FieldRuleMappingMode.FTW_MAPPED,
            priority="MEDIUM",
            source="Schedule A",
            form_section="Schedule A - Discovered FTW fields",
            field_type="Dynamic",
            existing_behavior="Review Only",
            new_behavior="Keep FTW",
            aliases=["Health", "Medical coverage"],
        )
        vision_rule = health_rule.model_copy(
            update={
                "key": "ftw_discovered_schedule_a_vision_ind",
                "label": "Vision Indicator",
                "ftw_field": "Vision Indicator",
                "xml_tag": "VisionInd",
                "aliases": ["Vision", "Eye care"],
            }
        )

        fields = extract_fields_from_document_text(
            b"Benefits: Health, Dental, Vision, Prescription Drug",
            "schedule-a.txt",
            rules=[health_rule, vision_rule],
        )
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(by_name["Health Indicator"], "Yes")
        self.assertEqual(by_name["Vision Indicator"], "Yes")

    def test_local_parser_keeps_one_best_value_per_schedule_a_field(self):
        text = """
        Name of Insurance Carrier: Kaiser Foundation Health Plan, Inc.
        Total Amount of Commissions Paid: $4,810.38
        March 12, 2026
        """

        fields = extract_fields_from_document_text(text.encode(), "schedule-a.txt")
        names = [field.field_name for field in fields]

        self.assertEqual(names.count("1a. Name of Insurance Company"), 1)
        self.assertEqual(names.count("3b. Amount of Commissions"), 1)

    def test_merge_replaces_invalid_ai_contract_with_validated_document_value(self):
        ai_fields = [
            NormalizedExtractionField(
                field_name="1d. Contract/Policy Number",
                value="4",
                confidence=0.99,
                page=1,
            ),
            NormalizedExtractionField(
                field_name="1d. Contract/Policy Number",
                value="10420761002",
                confidence=0.98,
                page=1,
            ),
        ]
        document_fields = [
            NormalizedExtractionField(
                field_name="1d. Contract/Policy Number",
                value="1042075/6-1001/1002",
                confidence=0.92,
                page=1,
            )
        ]

        merged = merge_schedule_a_fields(ai_fields, document_fields)

        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0].value, "1042075/6-1001/1002")
        self.assertEqual(
            merged[0].candidate_values,
            ["10420761002", "1042075/6-1001/1002"],
        )

    def test_groundx_query_contains_published_aliases_for_the_relevant_form(self):
        custom_alias = "Carrier Registry Number"
        rules = [
            rule.model_copy(update={"aliases": [*rule.aliases, custom_alias]})
            if rule.key == "schedule_a_part_i_1c_naic_code"
            else rule
            for rule in DEFAULT_FIELD_RULES
        ]

        query = build_groundx_schema_query(
            "schedule-a.pdf",
            rules,
            form_type=FormType.SCHEDULE_A,
        )

        self.assertIn(custom_alias, query)
        self.assertNotIn("Plan Administrator Name", query)

    def test_groundx_query_contains_alias_for_new_discovered_ftw_field(self):
        rule = FieldRule(
            key="ftw_discovered_schedule_a_ins_fail_provide_info_text",
            label="Reason information was not provided",
            ftw_field="Insurance Carrier Missing Information Explanation",
            xml_tag="InsFailProvideInfoText",
            mapping_mode=FieldRuleMappingMode.FTW_MAPPED,
            priority="MEDIUM",
            source="Schedule A",
            form_section="Schedule A - Discovered FTW fields",
            field_type="Dynamic",
            existing_behavior="Review Only",
            new_behavior="Keep FTW",
            aliases=["Carrier explanation for missing information"],
        )

        query = build_groundx_schema_query(
            "schedule-a.pdf",
            [rule],
            form_type=FormType.SCHEDULE_A,
        )

        self.assertIn("Reason information was not provided", query)
        self.assertIn("Carrier explanation for missing information", query)

    def test_groundx_xray_maps_gross_premium_table_to_nonexperience_line_10a(self):
        payload = {
            "chunks": [
                {
                    "pageNumbers": [2],
                    "text": "Total Premium received - Type of Benefit - Gross Premium",
                    "json": [
                        {
                            "record_type": "total_premium_received",
                            "rows": [
                                {"type_of_benefit": "Dental", "gross_premium": "$95,409.74"},
                                {"type_of_benefit": "Total", "gross_premium": "$95,409.74"},
                            ],
                        }
                    ],
                }
            ],
            "documentPages": [{"pageNumber": 2}],
        }

        fields = extract_fields_from_groundx_xray(payload)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(
            by_name["10a. Total premiums or subscription charges paid to carrier"],
            "95,409.74",
        )

    def test_groundx_ocr_maps_vendor_total_premium_to_nonexperience_line_10a(self):
        text = """
        5500 Schedule A Insurance Information
        Name of insurance carrier Sun Life Assurance Company of Canada
        Total Premium received 01/01/2025 to 12/31/2025
        Type of Benefit Dental Gross Premium $95,409.74
        Total $95,409.74
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(
            by_name["10a. Total premiums or subscription charges paid to carrier"],
            "95,409.74",
        )
        mapped = map_extraction_to_rules(
            "peerless-test",
            fields,
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )["fields"]
        self.assertEqual(classify_schedule_a_fields(mapped).contract_type.value, "NONEXPERIENCE_RATED")

    def test_groundx_ocr_maps_policy_year_premium_sentence_to_nonexperience_line_10a(self):
        text = """
        Annual Policy Information Report
        Premiums, Commissions and Fees are as paid during the policy year.
        Total premiums paid to Insurance Company during the policy year: $12,013.49
        See below for total commissions and fees paid by Insurance Company.
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(
            by_name["10a. Total premiums or subscription charges paid to carrier"],
            "12,013.49",
        )

    def test_annual_policy_report_extracts_colon_labelled_policy_year_dates(self):
        text = """
        Annual Policy Information Report
        Name of Insurance Carrier
        Life Insurance Company of North America
        EIN 23-1503749
        NAIC Code 65498
        Contract/Policy Number FLX0966852
        Contract/Policy Year From: 01/01/2025
        Contract/Policy Year To: 12/31/2025
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "12/31/2025")

    def test_groundx_ocr_maps_nonparticipating_subscription_charge_wording_to_line_10a(self):
        text = """
        Official ERISA Notification
        7. NON-PARTICIPATING CONTRACTS (PREMIUMS):
        (A) TOTAL PREMIUM OR SUBSCRIPTION CHARGES PAID
        TO CARRIER. $49,429.65
        (B) PREMIUMS DUE AND UNPAID AT END OF THE PLAN YEAR. $.00
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(
            by_name["10a. Total premiums or subscription charges paid to carrier"],
            "49,429.65",
        )
        mapped = map_extraction_to_rules(
            "oxford-test",
            fields,
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )["fields"]
        self.assertEqual(classify_schedule_a_fields(mapped).contract_type.value, "NONEXPERIENCE_RATED")

    def test_groundx_ocr_maps_experience_section_amounts_to_line_9(self):
        text = """
        Schedule A Part III
        9 Experience-rated contracts
        9a. Premiums: (1) Amount Received $1,739,422
        9b(1). Benefit Charges (1) Claims paid $1,503,774
        9c(1)(H). Total retention $235,648
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(by_name["9a. Premiums: (1) Amount Received"], "1,739,422")
        self.assertEqual(by_name["9b(1). Benefit Charges (1) Claims paid"], "1,503,774")
        self.assertEqual(by_name["9c(1)(H). Total retention"], "235,648")
        mapped = map_extraction_to_rules(
            "experience-test",
            fields,
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )["fields"]
        self.assertEqual(classify_schedule_a_fields(mapped).contract_type.value, "EXPERIENCE_RATED")

    def test_schedule_a_parser_extracts_fee_org_code_and_derived_purpose(self):
        text = """
        SCHEDULE A (Form 5500) 2024
        A Name of plan MIDWEST HOSE AND SPECIALTY HEALTH AND WELFARE BENEFITS PLAN
        B Three-digit plan number (PN) 501
        C Plan sponsor's name as shown on line 2a of Form 5500
        MIDWEST HOSE & SPECIALTY INC.
        D Employer Identification Number (EIN) 73-1185740

        (a) Name of insurance carrier
        UNITEDHEALTHCARE INSURANCE COMPANY
        (b) EIN 36-2739571
        (c) NAIC code 79413
        (d) Contract or identification number 1246876
        (e) Approximate number of persons covered at end of policy or contract year * 61
        (f) From 10/01/2024
        (g) To 09/30/2025

        fees paid/amount:  $0.00
        (a) Name and address of the agents, brokers or other persons to whom commissions or fees were paid:
        NFP CORPORATION SERVICES (OK) LLC
        4811 GAILLARDIA PKWY STE 300
        OKLAHOMA CITY OK 73142-1875
        (b) Amount of commissions paid: $111,892.96
        (c) Fees paid / Amount: $0.00
        (d) Purpose: N/A
        (e) Organizational Code: 3

        10a. Total premiums or subscription charges paid to carrier: $3,102,445.38
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(by_name["1a. Name of Insurance Company"], "UNITEDHEALTHCARE INSURANCE COMPANY")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "36-2739571")
        self.assertEqual(by_name["1c. NAIC Code"], "79413")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "1246876")
        self.assertEqual(by_name["3a. Name of Agent/Broker/Person"], "NFP CORPORATION SERVICES (OK) LLC")
        self.assertEqual(by_name["3b. Amount of Commissions"], "111,892.96")
        self.assertEqual(by_name["3c. Amount of Fees"], "0.00")
        self.assertEqual(by_name["3d. Purpose"], "COMMISSIONS")
        self.assertEqual(by_name["3e. Organizational Code"], "3")
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "3,102,445.38")

    def test_schedule_a_parser_extracts_kaiser_carrier_specific_labels(self):
        text = """
        INSURANCE INFORMATION
        Part I: Information Concerning Insurance Coverage, Fees, and Commissions
        Name of Insurance Carrier: Kaiser Foundation Health Plan, Inc.
        Plan Sponsor's Name: SPECIAL SERVICE FOR GROUPS, INC.
        Information Concerning Insurance Contract Coverage
        Kaiser Foundation Health Plan Region: CA
        Insurance Carrier: Kaiser Foundation Health Plan, Inc.
        Insurance Carrier Employer Identification Number: 94-1340523
        Insurance Carrier NAIC Code: 00000
        Plan Sponsor Contract or Identification Number: 608066
        Approximate number of persons covered at end of policy contract year: 37
        Contract Year from 01/2025 - 12/2025
        Information Concerning Insurance Contract Fees and Commissions
        Total Amount of Commissions Paid: $4,810.38
        Total Amount of Fees Paid: $0.00
        1) Name and address of the agent, broker, or other person to whom commissions or fees were paid:
        Gallagher Benefit Services, Inc.
        500 N BRAND BLVD STE 100
        GLENDALE, CA 91203-3931
        Amount of sales and base commissions paid to Gallagher Benefit Services, Inc.: $4,810.38
        Fees and other compensation paid to Gallagher Benefit Services, Inc.: $0.00
        Part III: Welfare Benefit Contract Information
        Premium applied by Kaiser Foundation Health Plan, Inc. during your plan's contract year: $262,735.11
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Kaiser Foundation Health Plan, Inc.")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "94-1340523")
        self.assertEqual(by_name["1c. NAIC Code"], "00000")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "608066")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "37")
        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(by_name["3a. Name of Agent/Broker/Person"], "Gallagher Benefit Services, Inc.")
        self.assertEqual(by_name["3b. Amount of Commissions"], "4,810.38")
        self.assertEqual(by_name["3c. Amount of Fees"], "0.00")
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "262,735.11")

    def test_schedule_a_parser_extracts_aetna_table_values_without_swapping_columns(self):
        text = """
        INSURANCE INFORMATION
        AETNA LIFE INSURANCE COMPANY AND AFFILIATES
        For Fiscal Plan Year beginning 01/01/2025 and ending 12/31/2025
        PART I Information Concerning Insurance Contract Coverage, Fees, and Commissions.
        1. Coverage: HNO Prospective
        (a) Name of Insurance Carrier:
        Aetna Health, Inc.
        (b) EIN: See Attached
        (c) NAIC Code: See Attached Listing
        (d) Contract Number
        or Identification:
        0847233HNO
        (e) Approximate Number of
        persons covered at the end
        of policy or contract year:
        610
        Policy or contract Year
        (f) From:
        01/01/2025
        (g) To:
        12/31/2025
        2. Insurance Fees and commissions paid to agents and brokers:
        Contract or Identification
        (a) Name and address of the agents or brokers to whom commissions or fees were paid.
        (b) Amount of commissions paid
        (c) & (d) Fees Paid Amount Purpose
        0847233HNO
        GALLAGHER BENEFIT SERVICES INC 505 N BRAND BLVD GLENDALE, CA 91203
        $47,063.49
        """

        fields = parse_schedule_a_text(text)
        by_name = {field.field_name: field.value for field in fields}

        self.assertEqual(by_name["1d. Contract/Policy Number"], "0847233HNO")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "610")
        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(by_name["3b. Amount of Commissions"], "47,063.49")

        mapped = map_extraction_to_rules(
            "test-filing",
            fields,
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )["fields"]
        mapped_by_label = {field.mapped_label: field.proposed_value for field in mapped}
        self.assertEqual(mapped_by_label["1d. Contract/Policy Number"], "0847233HNO")
        self.assertNotEqual(mapped_by_label["1d. Contract/Policy Number"], "610")

    def test_schedule_a_parser_prioritizes_metlife_broker_totals_over_page_fragments(self):
        pages = [
            """
            Schedule A
            Insurance Information
            Cover Letter
            Customer Name: THE ADVERTISING COUNCIL, INC.
            Attention: NANCY WING

            METROPOLITAN LIFE INSURANCE COMPANY
            EIN NAIC Code Contract or identification # Approximate number of persons covered at end of policy or contract year Policy or contract year
            13-5581829 65978 5955240 255 01/01/2025 12/31/2025
            """,
            """
            Totals
            Total Amount of commissions paid:1,998 Total fees paid/amount:345
            """,
            """
            Name and address of the agents, brokers or other persons to whom commissions or fees were paid
            Name: NFP CORPORATE SERVICES NY LLCAddress Line 1:PO BOX 9101
            Zip Code: 11803-9001 Organization
            code: 03

            Commissions Paid
            Coverage Amount Purpose
            Vision 1,576Base Commissions
            1,576Sub Total

            Fees Paid
            Coverage Amount Purpose
            Multiple 44 Non-Monetary
            Compensation
            44 Sub Total
            """,
            """
            Zip Code: 10166-3201 Organization
            code: 03
            Commissions Paid
            Coverage Amount Purpose
            0 Sub Total
            Fees Paid
            Coverage Amount Purpose
            Vision 2 Marketing Fees
            2 Sub Total

            Part III
            8. Experience-rated contracts
            N/A
            9. Nonexperience-rated contracts
            a. Total premiums or subscription charges paid to carrier:
            Vision 15,870
            """,
        ]
        fields = []
        for page, text in enumerate(pages, start=1):
            fields.extend(parse_schedule_a_text(text, page))
        fields.extend(parse_schedule_a_text("\n\n".join(pages), None))

        mapped = map_extraction_to_rules(
            "test-filing",
            fields,
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )["fields"]
        mapped_by_label = {field.mapped_label: field.proposed_value for field in mapped}

        self.assertEqual(mapped_by_label["3a. Name of Agent/Broker/Person"], "NFP CORPORATE SERVICES NY LLC")
        self.assertEqual(mapped_by_label["1d. Contract/Policy Number"], "5955240")
        self.assertNotEqual(mapped_by_label["1d. Contract/Policy Number"], "576Ba")
        self.assertEqual(mapped_by_label["3b. Amount of Commissions"], "1,998")
        self.assertEqual(mapped_by_label["3c. Amount of Fees"], "345")
        self.assertEqual(mapped_by_label["3d. Purpose"], "COMMISSIONS & FEES")
        self.assertEqual(mapped_by_label["3e. Organizational Code"], "03")
        self.assertEqual(mapped_by_label["10a. Total premiums or subscription charges paid to carrier"], "15,870")
        self.assertEqual(
            mapped_by_label["11. Did the insurance company fail to provide any information necessary to complete Schedule A?"],
            "No",
        )

    def test_schedule_a_parser_extracts_metlife_repeatable_broker_rows(self):
        text = """
        Name and address of the agents, brokers or other persons to whom commissions or fees were paid
        Name: NFP CORPORATE SERVICES NY LLCAddress Line 1:PO BOX 9101
        City: PLAINVIEW State: NY
        Zip Code: 11803-9001 Organization
        code: 03
        Commissions Paid
        Coverage Amount Purpose
        Vision 1,576Base Commissions
        1,576Sub Total
        Fees Paid
        Coverage Amount Purpose
        Multiple 44 Non-Monetary
        Compensation
        44 Sub Total

        Name and address of the agents, brokers or other persons to whom commissions or fees were paid
        Name: NFP INS SERVICES INC Address Line 1: 1250 S CAPITAL OF TEXAS HWY
        Address Line 2: BLDG 2 STE 125 City: AUSTIN State: TX
        Zip Code: 78746-6446 Organization
        code: 03
        Commissions Paid
        Coverage Amount Purpose
        Dental 289Base Commissions
        Vision 133Base Commissions
        422Sub Total
        Fees Paid
        Coverage Amount Purpose
        0 Sub Total

        Name and address of the agents, brokers or other persons to whom commissions or fees were paid
        Name: NFP CORPORATE SERVICES NY LLC Address Line 1: 200 PARK AVE RM 3202
        Address Line 2: ATTN ACCOUNTING City: NEW YORK State: NY
        Zip Code: 10166-3201 Organization
        code: 03
        Commissions Paid
        Coverage Amount Purpose
        0 Sub Total
        Fees Paid
        Coverage Amount Purpose
        Vision 299Supplemental
        Compensation
        299 Sub Total

        Name and address of the agents, brokers or other persons to whom commissions or fees were paid
        Name: NFP CORPORATE SERVICES LLC Address Line 1: 200 PARK AVE RM 3202
        Address Line 2: ATTN ACCOUNTING City: NEW YORK State: NY
        Zip Code: 10166-3201 Organization
        code: 03
        Commissions Paid
        Coverage Amount Purpose
        0 Sub Total
        Fees Paid
        Coverage Amount Purpose
        Vision 2 Marketing Fees
        2 Sub Total
        """

        rows = extract_schedule_a_broker_rows(text)

        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0].name, "NFP CORPORATE SERVICES NY LLC")
        self.assertEqual(rows[0].commission_total, "1,576")
        self.assertEqual(rows[0].fee_total, "44")
        self.assertEqual(rows[1].name, "NFP INS SERVICES INC")
        self.assertEqual(rows[1].commission_total, "422")
        self.assertEqual(rows[1].fee_total, "0")
        self.assertEqual(rows[2].fee_total, "299")
        self.assertEqual(rows[3].fee_total, "2")

    def test_schedule_a_parser_reads_metlife_address_and_split_state_zip(self):
        text = """
        (a) Name and address of the agents, brokers or other persons to whom commissions or fees were paid
        Name: GALLAGHER BENEFIT SERVICES INC
        Address: PO BOX 3009 City: ARLINGTON
        HEIGHTS ST: IL ZIP: 60006-3009
        Commissions Paid Fees Paid Organization
        codeCoverage Amount Purpose Coverage Amount Purpose
        Health 978 Base Commissions Multiple 107 Non-Monetary
        Compensation 03
        978 Sub-total 107 Sub-total
        """
        rows = extract_schedule_a_broker_rows(text)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].address_line_1, "PO BOX 3009")
        self.assertEqual(rows[0].city, "ARLINGTON HEIGHTS")
        self.assertEqual(rows[0].state, "IL")
        self.assertEqual(rows[0].zip_code, "60006-3009")
        self.assertEqual(rows[0].organization_code, "03")
        self.assertEqual(rows[0].commission_total, "978")
        self.assertEqual(rows[0].fee_total, "107")

    def test_broker_fragment_joins_only_same_name_and_zip(self):
        rows = merge_schedule_a_broker_rows(
            [
                ScheduleABrokerRow(name="GALLAGHER BENEFIT SERVICES INC", address_line_1="PO BOX 3009", city="ARLINGTON HEIGHTS", state="IL", zip_code="60006-3009", commission_total="978", fee_total="107"),
                ScheduleABrokerRow(name="GALLAGHER BENEFIT SERVICES INC", address_line_1="PO BOX 95287", city="CHICAGO", state="IL", zip_code="60690-7219", commission_total="0", fee_total="121"),
            ],
            [ScheduleABrokerRow(name="GALLAGHER BENEFIT SERVICES INC", city="ARLINGTON HEIGHTS ST: IL ZIP: 60006-3009", fee_total="3")],
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].address_line_1, "PO BOX 3009")
        self.assertEqual(rows[1].address_line_1, "PO BOX 95287")

    def test_broker_rows_same_recipient_remain_separate_by_coverage(self):
        address = {
            "name": "NFP CORPORATE SERVICES (NY) LLC",
            "address_line_1": "200 PARK AVE",
            "city": "NEW YORK",
            "state": "NY",
            "zip_code": "10166",
        }
        rows = merge_schedule_a_broker_rows(
            [
                ScheduleABrokerRow(
                    **address,
                    fee_total="48,230.78",
                )
            ],
            [
                ScheduleABrokerRow(
                    **address,
                    fee_total="48,230.78",
                    fee_rows=[
                        ScheduleABrokerMoneyRow(
                            coverage="LIFE INSURANCE",
                            amount="48,230.78",
                            purpose="Contingent Compensation",
                        )
                    ],
                ),
                ScheduleABrokerRow(
                    **address,
                    fee_total="26,143.92",
                    fee_rows=[
                        ScheduleABrokerMoneyRow(
                            coverage="LONG TERM DISABILITY",
                            amount="26,143.92",
                            purpose="Contingent Compensation",
                        )
                    ],
                ),
            ],
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual([row.fee_total for row in rows], ["48,230.78", "26,143.92"])

        service = FTWilliamsReviewService()
        life_rows = service._broker_rows_for_schedule_desc(rows, "Schedule A-LIFE")
        ltd_rows = service._broker_rows_for_schedule_desc(rows, "Schedule A-LTD")

        self.assertEqual([row.fee_total for row in life_rows], ["48,230.78"])
        self.assertEqual([row.fee_total for row in ltd_rows], ["26,143.92"])

    def test_broker_rows_are_filtered_by_contract_before_generated_schedule_description(self):
        rows = [
            ScheduleABrokerRow(
                name="HUMMEL GROUP",
                commission_total="3,406.53",
                commission_rows=[
                    ScheduleABrokerMoneyRow(
                        coverage="10049061001",
                        amount="3,406.53",
                        purpose="Commissions",
                    )
                ],
            ),
            ScheduleABrokerRow(
                name="HUMMEL GROUP",
                commission_total="245.62",
                commission_rows=[
                    ScheduleABrokerMoneyRow(
                        coverage="10049071001",
                        amount="245.62",
                        purpose="Commissions",
                    )
                ],
            ),
        ]

        selected = FTWilliamsReviewService()._broker_rows_for_schedule_desc(
            rows,
            "7-1",
            "10049071001",
        )

        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0].commission_total, "245.62")

    def test_schedule_a_parser_stops_last_broker_before_part_iii(self):
        text = """
        Name and address of the agents, brokers or other persons to whom commissions or fees were paid
        Name: NFP CORPORATE SERVICES LLC Address Line 1: 200 PARK AVE RM 3202
        Address Line 2: ATTN ACCOUNTING City: NEW YORK State: NY
        Zip Code: 10166-3201 Organization code: 03
        Commissions Paid
        Coverage Amount Purpose
        0 Sub Total
        Fees Paid
        Coverage Amount Purpose
        Vision 2 Marketing Fees
        2 Sub Total
        Part III Welfare Benefit Contract Information
        Vision 15,870
        If more than one contract covers the same group of employees, complete the information below.
        """

        rows = extract_schedule_a_broker_rows(text)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].fee_total, "2")
        self.assertEqual(rows[0].fee_rows, [])
        self.assertNotIn("15,870", str(rows[0].model_dump()))
        self.assertNotIn("Welfare Benefit Contract Information", str(rows[0].model_dump()))

    def test_schedule_a_parser_extracts_bcbsma_worksheet_experience_rated_fields(self):
        text = """
        ACCOUNT NAME: R. H. White Construction Co. I
        ACCOUNT #: 0123307
        PERIOD: 01/01/2025 - 12/31/2025 @ 03/31/2026
        NAIC CODE: 53228
        EIN CODE: 04-1045815
        MEDICAL DENTAL SENIOR
        LAST MONTH OF PERIOD ENROLLMENT
        Employees 361 0 0
        Employee & Dependents 773 0 0
        PREMIUM
        Total Premium $6,554,192 $0 $0
        BENEFIT CHARGES
        Incurred Claims $5,498,349 $0 $0
        Incurred But Not Reported $44,878 $0 $0
        Claims Charged $5,543,227 $0 $0
        RETENTION ALLOCATION
        Base Commission $94,601 $0 $0
        Taxes $185 $0 $0
        Other Retention Charges $916,179 $0 $0
        Blue Cross Blue Shield of Massachusetts, Inc.
        FULLY INSURED #5500A WORKSHEET
        """

        fields = extract_bcbsma_schedule_a_worksheet_fields(text, page=1)
        by_name = {field.field_name: field.value for field in fields}
        summaries = extract_bcbsma_schedule_a_worksheet_summaries(text, page=1)

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Blue Cross Blue Shield of Massachusetts, Inc.")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "04-1045815")
        self.assertEqual(by_name["1c. NAIC Code"], "53228")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "0123307")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "773")
        self.assertEqual(by_name["9a. Premiums: (1) Amount Received"], "6,554,192")
        self.assertEqual(by_name["9b(1). Benefit Charges (1) Claims paid"], "5,498,349")
        self.assertEqual(by_name["9b(2). Increase (decrease) in claim reserves"], "44,878")
        self.assertEqual(by_name["9b(3). Incurred claims (add(1) and (2))"], "5,543,227")
        self.assertEqual(by_name["9c(1)(A). Commissions"], "94,601")
        self.assertEqual(by_name["9c(1)(E). Taxes"], "185")
        self.assertEqual(by_name["9c(1)(G). Other retention charges"], "916,179")
        self.assertEqual(by_name["9c(1)(H). Total retention"], "1,010,965")
        self.assertEqual(len(summaries), 1)
        self.assertEqual(summaries[0].source, "BCBSMA #5500A worksheet")
        self.assertEqual(summaries[0].account_number, "0123307")

    def test_schedule_a_parser_extracts_bcbsma_commission_breakdown_broker_row(self):
        text = """
        ACCOUNT NAME: R. H. White Construction Co. I
        ACCOUNT #: 0123307
        PERIOD: 01/01/2025 - 12/31/2025 @ 03/31/2026
        NAIC CODE: 53228
        EIN CODE: 04-1045815
        MEDICAL DENTAL SENIOR #VALUE!
        COMMISSION BREAKDOWN
        R S C INS BKGE DBA RISK STRATAGIES CO. $94,601.06 $8,932.00 $0.00
        ### OTHER COMMISSION *
        R S C INS BKGE DBA RISK STRATAGIES CO. $23,270.00
        ### NON MONETARY C0MPENSATION *
        R S C INS BKGE DBA RISK STRATAGIES CO. $209.15
        COMMISSIONS AND BONUS BREAKDOWN
        """

        rows = extract_bcbsma_commission_breakdown_broker_rows(text, page=3)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "R S C INS BKGE DBA RISK STRATAGIES CO.")
        self.assertEqual(rows[0].commission_total, "126,803")
        self.assertEqual(rows[0].fee_total, "209")
        self.assertEqual(rows[0].fee_rows[0].purpose, "Non-Monetary Compensation")

    def test_schedule_a_parser_extracts_bcbs_michigan_experience_addendum_without_10a_leak(self):
        pages = [
            (
                2,
                """
                SCHEDULE A (ERISA FORM 5500)
                INSURANCE INFORMATION
                GROUP NAME: OptimizeRX Corporation
                PART I: Insurance Information
                1. COVERAGE INFORMATION
                (a) NAME OF INSURANCE CARRIER BLUE CROSS BLUE SHIELD OF MICHIGAN
                (b) EMPLOYER IDENTIFICATION NUMBER (EIN) 38-2069753
                (c) NATIONAL ASSOCIATION OF INSURANCE COMMISSIONERS (NAIC) CODE 54291
                (d) CONTRACT OR IDENTIFICATION NUMBER 616888
                (e) APPROX. NUMBER OF PERSONS COVERED 250
                (f) POLICY OR CONTRACT YEAR FROM 12/1/2024
                (g) POLICY OR CONTRACT YEAR TO 11/30/2025
                2. INSURANCE FEE AND COMMISSION INFORMATION (SEE SCHEDULE A ADDENDUM)
                3. PERSONS RECEIVING COMMISSIONS AND FEES (SEE SCHEDULE A ADDENDUM)
                PART II: INVESTMENT AND ANNUITY CONTRACT INFORMATION NOT APPLICABLE
                PART III: WELFARE BENEFIT CONTRACT INFORMATION
                9. EXPERIENCE-RATED CONTRACTS
                (a) PREMIUMS:
                (i) AMOUNT RECEIVED $1,739,422
                (ii) AND (iii) NOT APPLICABLE
                (iv) AMOUNT EARNED $1,739,422
                (b) BENEFIT CHARGES:
                (i) CLAIMS PAID $2,088,380
                (ii) INCREASE (DECREASE) IN CLAIM RESERVES $2,012
                (iii) INCURRED CLAIMS (ADD (i) AND (ii)) $2,090,392
                (iv) CLAIMS CHARGED (NET OF EXCESS CLAIMS) $2,032,510
                (c) REMAINDER OF PREMIUM
                (i) RETENTION CHARGES
                A. COMMISSIONS NOT APPLICABLE
                B. ADMINISTRATIVE SERVICE OR OTHER FEES $199,707
                C. OTHER SPECIFIC ACQUISITION COSTS $0
                D. OTHER EXPENSES (SUBSIDIES, ETC.) $0
                E. ESTIMATED TAXES, FEES AND ASSESSMENTS $21,924
                F. CHARGES FOR RISK OR OTHER CONTINGENCIES $60,433
                G. OTHER RETENTION CHARGES (POOLING CHARGE) $178,752
                H. TOTAL RETENTION $460,817
                (ii) DIVIDENDS OR RETROACTIVE RATE REFUNDS (CREDITED) $0
                (d) STATUS OF POLICYHOLDER RESERVES AT END OF YEAR
                (i) AMOUNT HELD TO PROVIDE BENEFITS AFTER RETIREMENT NOT APPLICABLE
                (ii) CLAIMS RESERVES $109,724
                (iii) OTHER RESERVES $0
                (e) DIVIDENDS OR RETROACTIVE RATE REFUNDS DUE $0
                10. NONEXPERIENCE-RATED CONTRACTS NOT APPLICABLE
                PART IV: PROVISION OF INFORMATION
                """,
            ),
            (
                3,
                """
                Client Name: OPTIMIZERX CORPORATION
                Group Number: 007049950
                CID: 616888
                Contract Year From: 12/01/2024
                Contract Year To: 11/30/2025
                AGENT/BROKER COMMISSION & INCENTIVE PAYMENTS
                 -- Name and address of agent or broker: NFP CORPORATE SERVICES NY LLC
                340 MADISON AVE 21ST FLOOR
                NEW YORK, NY 10173-0173
                 -- Amount of Sales and Base Commissions Paid $0.00
                 -- Fees and Other Commissions Paid Amount $1,255.50
                 -- Non-Monetary Compensations to Plan
                 (gifts, meals, entertainments, etc.) $0.00
                 -- Organization Code (for Schedule (A) 3
                AGENT/BROKER COMMISSION & INCENTIVE PAYMENTS
                 -- Name and address of agent or broker: KATHERINE HENRY
                Nfp Corporate Services (ny), Llc 340 Madison Avenue
                New York, NY -
                 -- Amount of Sales and Base Commissions Paid $53,816.53
                 -- Fees and Other Commissions Paid Amount $0.00
                 -- Non-Monetary Compensations to Plan
                 (gifts, meals, entertainments, etc.) $0.00
                 -- Organization Code ( for Schedule (A) 3
                Blue Cross Blue Shield Michigan
                ADDENDUM TO SCHEDULE A/C (ERISA FORM 5500)
                """,
            ),
        ]

        fields = extract_bcbs_michigan_schedule_a_fields(pages)
        by_name = {field.field_name: field.value for field in fields}
        rows = extract_bcbs_michigan_addendum_broker_rows("\n".join(text for _, text in pages))
        summaries = extract_bcbs_michigan_schedule_a_summaries(pages)

        self.assertEqual(by_name["1a. Name of Insurance Company"], "BLUE CROSS BLUE SHIELD OF MICHIGAN")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "616888")
        self.assertEqual(by_name["9a. Premiums: (1) Amount Received"], "1,739,422")
        self.assertEqual(by_name["9b(1). Benefit Charges (1) Claims paid"], "2,088,380")
        self.assertEqual(by_name["9c(1)(B). Administrative service or other fees"], "199,707")
        self.assertEqual(by_name["9c(1)(H). Total retention"], "460,817")
        self.assertEqual(by_name["9d(2). Claim reserves"], "109,724")
        self.assertEqual(by_name["3b. Amount of Commissions"], "53,816.53")
        self.assertEqual(by_name["3c. Amount of Fees"], "1,255.50")
        self.assertNotIn("10a. Total premiums or subscription charges paid to carrier", by_name)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].name, "NFP CORPORATE SERVICES NY LLC")
        self.assertEqual(rows[0].fee_total, "1,255.50")
        self.assertEqual(rows[1].name, "KATHERINE HENRY")
        self.assertEqual(rows[1].commission_total, "53,816.53")
        self.assertEqual(summaries[0].source, "BCBS Michigan Schedule A/C addendum")

    def test_schedule_a_parser_extracts_eyemed_worksheet_as_combined_ftw_record(self):
        pages = [
            (
                1,
                """
                Vision Insurance Information For Form 5500
                Report Start DateReport End Date
                1/1/25 12/31/25
                Information Compiled By: EyeMed Vision Care on behalf of the Fidelity Security Life Insurance Company
                Name of Plan Contract orID # Enrollment Group
                Approximate number ofsubscribers covered atend of policy or contractyear:
                Approximate number ofsubscribers anddependents covered at endof policy or contract year:EIN NAIC Amount
                RH WHITE CONSTRUCTION COMPANIES10258481001
                RH WHITECONSTRUCTIONCOMPANIES 175 366 43094984471870 $23,875.21
                RH WHITE CONSTRUCTION COMPANIESCOBRA 10258491001
                RH WHITECONSTRUCTIONCOMPANIES COBRA 1 1 43094984471870 $315.31
                176 367 Total: $24,190.52
                Payee Name Contract or ID # Address Line 1 City StateZip Code Amount
                RSC Insurance Brokerage, Inc. - Boston,10258481001160 Federal Street Boston MA 02110 $2,979.55
                RSC Insurance Brokerage, Inc. - Boston,10258491001160 Federal Street Boston MA 02110 $32.42
                Selman & Company, LLC 10258481001One Integrity Parkway Cleveland OH 44143 $1,845.07
                Selman & Company, LLC 10258491001One Integrity Parkway Cleveland OH 44143 $19.50
                Total: $4,876.54
                Commissions or fees paid by carrier to agents, brokers or other persons:
                Payments Received by carrier from plan or plan sponsor:
                """,
            )
        ]

        fields = extract_eyemed_schedule_a_fields(pages)
        by_name = {field.field_name: field.value for field in fields}
        summaries = extract_eyemed_schedule_a_summaries(pages)
        rows = extract_eyemed_broker_rows(pages)

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Fidelity Security Life Insurance Company")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "43-0949844")
        self.assertEqual(by_name["1c. NAIC Code"], "71870")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "1025848/9-1001")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "366")
        self.assertEqual(by_name["1f. Policy Year Beginning Date"], "01/01/2025")
        self.assertEqual(by_name["1g. Policy Year Ending Date"], "12/31/2025")
        self.assertEqual(by_name["3b. Amount of Commissions"], "4,876.54")
        self.assertEqual(by_name["3c. Amount of Fees"], "0")
        self.assertNotIn("3d. Purpose", by_name)
        self.assertNotIn("3e. Organizational Code", by_name)
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "24,190.52")
        self.assertTrue(all(field.page == 1 for field in fields))
        resolved = resolve_schedule_a_result(
            NormalizedExtractionResult(
                provider="EyeMed vision worksheet parser",
                fields=fields,
                schedule_a_broker_rows=rows,
                raw={},
            )
        )
        self.assertEqual(
            next(field for field in resolved.fields if field.field_name.startswith("3c.")).decision,
            "AUTOMATIC",
        )
        self.assertEqual(len(summaries), 1)
        self.assertEqual(summaries[0].source, "EyeMed vision worksheet")
        self.assertEqual(len(summaries[0].benefit_rows), 2)
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0].name, "RSC Insurance Brokerage, Inc. - Boston")
        self.assertEqual(rows[0].address_line_1, "160 Federal Street")
        self.assertEqual(rows[0].city, "Boston")
        self.assertEqual(rows[0].commission_total, "2,979.55")
        self.assertEqual(rows[-1].commission_total, "19.50")

    def test_schedule_a_parser_extracts_eyemed_rows_with_blank_identifier_cells(self):
        pages = [
            (
                1,
                """
                Vision Insurance Information For Form 5500
                Information Compiled By: EyeMed Vision Care on behalf of the Fidelity Security Life Insurance Company
                Report Start DateReport End Date
                1/1/2025 12/31/2025
                Name of Plan Contract orID # Enrollment Group
                Approximate number ofsubscribers covered atend of policy or contractyear:
                Approximate number ofsubscribers anddependents covered at endof policy or contract year:EIN NAIC Amount
                OXFORD BIOMEDICA (US) LLC10420751001OXFORD BIOMEDICA (US)LLC 134 335 43094984471870 $11,646.30
                OXFORD BIOMEDICA (US) LLC10420751002OXFORD BIOMEDICA US,INC. 0 0 $0.00
                OXFORD BIOMEDICA (US) LLC COBRA10420761001OXFORD BIOMEDICA (US)LLC COBRA 0 0 $94.40
                OXFORD BIOMEDICA (US) LLC COBRA10420761002OXFORD BIOMEDICA US,INC. COBRA 1 4 43094984471870 $224.48
                Total: $11,965.18
                Payee Name Contract or ID # Address Line 1 City StateZip Code Amount
                RSC Insurance Brokerage, Inc. - Boston,10420751001160 Federal Street Boston MA 02110 $4,870.24
                RSC Insurance Brokerage, Inc. - Boston,10420761001160 Federal Street Boston MA 02110 $18.88
                RSC Insurance Brokerage, Inc. - Boston,10420761002160 Federal Street Boston MA 02110 $31.05
                Total: $4,920.17
                Commissions or fees paid by carrier to agents, brokers or other persons:
                Payments Received by carrier from plan or plan sponsor:
                """,
            )
        ]

        fields = extract_eyemed_schedule_a_fields(pages)
        by_name = {field.field_name: field.value for field in fields}
        summaries = extract_eyemed_schedule_a_summaries(pages)
        rows = extract_eyemed_broker_rows(pages)

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Fidelity Security Life Insurance Company")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "43-0949844")
        self.assertEqual(by_name["1c. NAIC Code"], "71870")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "1042075/6-1001/1002")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "335")
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "11,965.18")
        self.assertEqual(by_name["3b. Amount of Commissions"], "4,920.17")
        self.assertEqual(len(summaries), 1)
        self.assertEqual(len(summaries[0].benefit_rows), 4)
        self.assertEqual(len(rows), 3)

    def test_eyemed_broker_rows_accept_city_joined_to_state_without_losing_cents(self):
        pages = [
            (
                1,
                """
                Vision Insurance Information For Form 5500
                Information Compiled By: EyeMed Vision Care on behalf of the Fidelity Security Life Insurance Company
                Payee Name Contract or ID # Address Line 1 City StateZip Code Amount
                NFP Corporate Services NY - Norwell, MA10545381001Po Box 786677 PhiladelphiaPA 19178-6677 $1,932.21
                Total: $1,932.21
                Commissions or fees paid by carrier to agents, brokers or other persons:
                """,
            )
        ]

        rows = extract_eyemed_broker_rows(pages)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "NFP Corporate Services NY - Norwell, MA")
        self.assertEqual(rows[0].city, "Philadelphia")
        self.assertEqual(rows[0].state, "PA")
        self.assertEqual(rows[0].zip_code, "19178-6677")
        self.assertEqual(rows[0].commission_total, "1,932.21")

    def test_groundx_failure_marks_local_fallback_for_manual_review(self):
        fallback = NormalizedExtractionResult(
            provider="Local PDF parser fallback",
            fields=[NormalizedExtractionField(field_name="1e. Persons Covered (End of Policy Year)", value="143740", confidence=0.99)],
            schedule_a_broker_rows=[ScheduleABrokerRow(name="NFP Corporate Services", confidence=0.95)],
        )
        settings = SimpleNamespace(groundx_api_key="test", groundx_bucket_id="test")
        service = ExtractionService()

        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.extract_schedule_a_classification_signals", return_value=[]),
            patch("app.services.extractor.local_schedule_a_pdf_result", return_value=fallback),
            patch.object(service, "_extract_with_groundx", new=AsyncMock(side_effect=RuntimeError("temporary AI failure"))),
        ):
            result = asyncio.run(service.extract_schedule_a(b"pdf", "schedule-a.pdf"))

        self.assertIn("manual review required", result.provider.lower())
        self.assertTrue(result.raw["manual_review_required"])
        self.assertLessEqual(result.fields[0].confidence, 0.5)
        self.assertLessEqual(result.schedule_a_broker_rows[0].confidence, 0.5)

    def test_groundx_stall_is_bounded_and_uses_local_fallback(self):
        fallback = NormalizedExtractionResult(
            provider="Local PDF parser fallback",
            fields=[
                NormalizedExtractionField(field_name="1a. Name of Insurance Company", value="TEST CARRIER", confidence=0.96),
                NormalizedExtractionField(field_name="1b. Insurance Carrier EIN", value="12-3456789", confidence=0.98),
                NormalizedExtractionField(field_name="1c. NAIC Code", value="12345", confidence=0.98),
                NormalizedExtractionField(field_name="1d. Contract/Policy Number", value="ABC123", confidence=0.97),
                NormalizedExtractionField(field_name="1e. Persons Covered (End of Policy Year)", value="10", confidence=0.97),
                NormalizedExtractionField(field_name="1f. Policy Year Beginning Date", value="01/01/2025", confidence=0.97),
                NormalizedExtractionField(field_name="1g. Policy Year Ending Date", value="12/31/2025", confidence=0.97),
            ],
        )
        settings = SimpleNamespace(
            groundx_api_key="test",
            groundx_bucket_id="test",
            groundx_max_wait_seconds=1,
        )

        async def stalled_groundx(*_args, **_kwargs):
            await asyncio.sleep(5)

        service = ExtractionService()
        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.extract_schedule_a_classification_signals", return_value=[]),
            patch("app.services.extractor.local_schedule_a_pdf_result", return_value=fallback),
            patch.object(service, "_extract_with_groundx", side_effect=stalled_groundx),
        ):
            result = asyncio.run(service.extract_schedule_a(b"pdf", "schedule-a.pdf"))

        self.assertIn("verified local fallback", result.provider.lower())
        self.assertIn("TimeoutError", result.raw["fallback_reason"])

    def test_groundx_document_lookup_is_scoped_to_configured_bucket(self):
        response = SimpleNamespace(
            status_code=200,
            json=lambda: {
                "documents": [
                    {
                        "documentId": "document-1",
                        "fileName": "schedule-a.pdf",
                        "status": "complete",
                        "xrayUrl": "https://upload.groundx.ai/xray.json",
                    }
                ]
            },
        )
        client = SimpleNamespace(get=AsyncMock(return_value=response))
        settings = SimpleNamespace(
            groundx_bucket_id=35071,
            groundx_max_wait_seconds=3,
            groundx_poll_seconds=1,
        )
        service = ExtractionService()

        with patch("app.services.extractor.get_settings", return_value=settings):
            refs = asyncio.run(
                service._find_groundx_document_refs(
                    client,
                    "https://api.groundx.ai/api/v1",
                    {"X-API-Key": "test"},
                    [{"processId": "process-1", "status": "complete"}],
                    "schedule-a.pdf",
                )
            )

        self.assertEqual(refs[0]["documentId"], "document-1")
        client.get.assert_awaited_once_with(
            "https://api.groundx.ai/api/v1/ingest/documents/35071",
            headers={"X-API-Key": "test"},
        )

    def test_groundx_poll_recognizes_nested_ingest_status(self):
        response = SimpleNamespace(
            json=lambda: {"ingest": {"processId": "process-1", "status": "complete"}},
            raise_for_status=lambda: None,
        )
        client = SimpleNamespace(get=AsyncMock(return_value=response))
        settings = SimpleNamespace(
            groundx_max_wait_seconds=3,
            groundx_poll_seconds=1,
        )
        service = ExtractionService()

        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.asyncio.sleep", new=AsyncMock()),
        ):
            result = asyncio.run(
                service._poll_groundx_process(
                    client,
                    "https://api.groundx.ai/api/v1",
                    {"X-API-Key": "test"},
                    "process-1",
                )
            )

        self.assertEqual(result["ingest"]["status"], "complete")
        self.assertEqual(client.get.await_count, 1)

    def test_groundx_xray_fetch_retries_when_artifact_is_not_ready(self):
        not_ready = SimpleNamespace(status_code=404)
        ready = SimpleNamespace(
            status_code=200,
            raise_for_status=lambda: None,
            json=lambda: {"chunks": [], "documentPages": []},
        )
        client = SimpleNamespace(get=AsyncMock(side_effect=[not_ready, ready]))
        settings = SimpleNamespace(
            groundx_xray_fetch_attempts=2,
            groundx_poll_seconds=1,
        )
        service = ExtractionService()

        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.asyncio.sleep", new=AsyncMock()) as sleep,
        ):
            result = asyncio.run(
                service._fetch_groundx_xray(
                    client,
                    "https://api.groundx.ai/api/v1",
                    {"X-API-Key": "test"},
                    "document-1",
                )
            )

        self.assertEqual(result, {"chunks": [], "documentPages": []})
        self.assertEqual(client.get.await_count, 2)
        sleep.assert_awaited_once_with(1)

    def test_schedule_a_uses_separate_groundx_operation_timeout(self):
        groundx = NormalizedExtractionResult(
            provider="GroundX structured extract",
            fields=[
                NormalizedExtractionField(
                    field_name="1a. Name of Insurance Company",
                    value="Test Carrier",
                    confidence=0.99,
                )
            ],
        )
        local = NormalizedExtractionResult(provider="Local PDF parser", fields=[])
        settings = SimpleNamespace(
            groundx_api_key="test",
            groundx_bucket_id="test",
            groundx_max_wait_seconds=90,
            groundx_operation_timeout_seconds=240,
        )
        observed: dict[str, float] = {}

        async def capture_timeout(awaitable, *, timeout):
            observed["timeout"] = timeout
            return await awaitable

        service = ExtractionService()
        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.extract_schedule_a_classification_signals", return_value=[]),
            patch("app.services.extractor.local_schedule_a_pdf_result", return_value=local),
            patch("app.services.extractor.asyncio.wait_for", side_effect=capture_timeout),
            patch.object(service, "_extract_with_groundx", new=AsyncMock(return_value=groundx)),
        ):
            result = asyncio.run(service._extract_schedule_a_unresolved(b"pdf", "schedule-a.pdf"))

        self.assertEqual(result.provider, "GroundX structured extract")
        self.assertEqual(observed["timeout"], 240)

    def test_groundx_failure_keeps_complete_local_fallback_trusted(self):
        fallback = NormalizedExtractionResult(
            provider="Local PDF parser fallback",
            fields=[
                NormalizedExtractionField(field_name="1a. Name of Insurance Company", value="Fidelity Security Life Insurance Company", confidence=0.96),
                NormalizedExtractionField(field_name="1b. Insurance Carrier EIN", value="43-0949844", confidence=0.98),
                NormalizedExtractionField(field_name="1c. NAIC Code", value="71870", confidence=0.98),
                NormalizedExtractionField(field_name="1d. Contract/Policy Number", value="1054538/9-1001", confidence=0.97),
                NormalizedExtractionField(field_name="1e. Persons Covered (End of Policy Year)", value="171", confidence=0.97),
                NormalizedExtractionField(field_name="1f. Policy Year Beginning Date", value="04/01/2025", confidence=0.97),
                NormalizedExtractionField(field_name="1g. Policy Year Ending Date", value="03/31/2026", confidence=0.97),
            ],
            schedule_a_broker_rows=[ScheduleABrokerRow(name="NFP Corporate Services", confidence=0.95)],
        )
        settings = SimpleNamespace(groundx_api_key="test", groundx_bucket_id="test")
        service = ExtractionService()

        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.extract_schedule_a_classification_signals", return_value=[]),
            patch("app.services.extractor.local_schedule_a_pdf_result", return_value=fallback),
            patch.object(service, "_extract_with_groundx", new=AsyncMock(side_effect=RuntimeError("temporary AI failure"))),
        ):
            result = asyncio.run(service.extract_schedule_a(b"pdf", "schedule-a.pdf"))

        self.assertIn("verified local fallback", result.provider.lower())
        self.assertTrue(result.raw["fallback_validated"])
        self.assertEqual(result.fields[0].confidence, 0.96)
        self.assertEqual(result.schedule_a_broker_rows[0].confidence, 0.95)

    def test_groundx_result_is_defensively_supplemented_by_local_parser(self):
        groundx = NormalizedExtractionResult(
            provider="GroundX X-Ray",
            fields=[NormalizedExtractionField(field_name="1a. Name of Insurance Company", value="Fidelity", confidence=0.5)],
        )
        local = NormalizedExtractionResult(
            provider="Local PDF parser",
            fields=[
                NormalizedExtractionField(field_name="1a. Name of Insurance Company", value="Fidelity Security Life Insurance Company", confidence=0.96),
                NormalizedExtractionField(field_name="1b. Insurance Carrier EIN", value="43-0949844", confidence=0.98),
            ],
            schedule_a_broker_rows=[ScheduleABrokerRow(name="NFP Corporate Services", confidence=0.95)],
        )
        settings = SimpleNamespace(groundx_api_key="test", groundx_bucket_id="test")
        service = ExtractionService()

        with (
            patch("app.services.extractor.get_settings", return_value=settings),
            patch("app.services.extractor.extract_schedule_a_classification_signals", return_value=[]),
            patch("app.services.extractor.local_schedule_a_pdf_result", return_value=local),
            patch.object(service, "_extract_with_groundx", new=AsyncMock(return_value=groundx)),
        ):
            result = asyncio.run(service.extract_schedule_a(b"pdf", "schedule-a.pdf"))

        by_name = {field.field_name: field.value for field in result.fields}
        self.assertEqual(by_name["1a. Name of Insurance Company"], "Fidelity Security Life Insurance Company")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "43-0949844")
        self.assertEqual(result.schedule_a_broker_rows[0].name, "NFP Corporate Services")

    def test_schedule_a_parser_extracts_standard_long_form_separate_benefits(self):
        pages = self._standard_long_form_pages()

        records = extract_standard_schedule_a_records(pages)
        summaries = extract_standard_schedule_a_summaries(pages)
        rows = extract_standard_broker_rows(pages)

        by_coverage = {record["coverage"]: record for record in records}
        self.assertEqual(set(by_coverage), {"DENTAL", "LIFE INSURANCE", "LONG TERM DISABILITY"})

        dental = by_coverage["DENTAL"]
        self.assertEqual(dental["carrier_name"], "Standard Insurance Company")
        self.assertEqual(dental["contract_number"], "168262")
        self.assertEqual(dental["persons_covered"], "63")
        self.assertEqual(dental["ein"], "93-0242990")
        self.assertEqual(dental["naic_code"], "69019")
        self.assertEqual(dental["commission_total"], "1,506.01")
        self.assertEqual(dental["fee_total"], "198.74")
        self.assertEqual(dental["experience_values"]["9a. Premiums: (1) Amount Received"], "30,312.84")
        self.assertEqual(dental["experience_values"]["9b(1). Benefit Charges (1) Claims paid"], "22,882.90")
        self.assertEqual(dental["experience_values"]["9c(1)(H). Total retention"], "9,627.22")

        life = by_coverage["LIFE INSURANCE"]
        self.assertEqual(life["persons_covered"], "107")
        self.assertEqual(life["commission_total"], "1,731.10")
        self.assertEqual(life["fee_total"], "236.46")
        self.assertEqual(life["experience_values"]["9a(3). Increase (decrease) in unearned premium reserve"], "-2,104.00")

        ltd = by_coverage["LONG TERM DISABILITY"]
        self.assertEqual(ltd["commission_total"], "1,308.98")
        self.assertEqual(ltd["fee_total"], "147.90")
        self.assertEqual(ltd["experience_values"]["9b(2). Increase (decrease) in claim reserves"], "1,610.73")

        self.assertEqual(len(summaries), 3)
        self.assertEqual([summary.coverage for summary in summaries], ["DENTAL", "LIFE INSURANCE", "LONG TERM DISABILITY"])
        self.assertEqual(len(rows), 3)
        self.assertEqual([row.commission_total for row in rows], ["1,506.01", "1,731.10", "1,308.98"])
        self.assertEqual([row.fee_total for row in rows], ["198.74", "236.46", "147.90"])
        self.assertEqual(rows[0].name, "LEAHY CONSULTING SERVICES")
        self.assertEqual(rows[0].organization_code, "3")

    def test_schedule_a_parser_extracts_standard_life_new_york_long_form(self):
        pages = [
            (
                1,
                """PAGE: 1
PART I
2) INSURANCE FEES AND COMMISSIONS PAID TO AGENTS, BROKERS AND OTHER PERSONS:
E) ORG.
CODE
NFP CORPORATE SERVICES (NY) LLC
200 PARK AVE 32ND FL
NEW YORK, NY 10166
$0.00 $48,230.78 $0.00 $0.00 3
Standard Life Ins Co of NY
NFP CORP
1/1/2025
12/31/2025
6,159
13-4119477
000-89009
$48,230.78
$0.00
THE FINANCIAL DATA BELOW IS PROVIDED FOR YOUR INFORMATION
LIFE INSURANCE
PLAN INFORMATION REPORT FOR THE PERIOD OF
753370
LONG FORM INFORMATION
1/1/2025
12/31/2025""",
            ),
            (
                2,
                """PART III -
EXPERIENCE RATED CONTRACTS
PLAN INFORMATION REPORT FOR THE PERIOD OF
Standard Life Ins Co of NY HEREBY CERTIFIES THAT THIS INFORMATION IS COMPLETE AND ACCURATE
($184,736.00)
$2,101.00
$2,683,562.19
$2,250,000.00
$230,723.00
$2,480,723.00
$2,480,723.00
$48,230.78
$0.00
$1,200.00
$251,041.00
$58,189.19
$151,268.05
$0.00
$509,929.02
$0.00
$0.00
$84,300.00
$0.00
$0.00
753370
LIFE INSURANCE
$2,870,399.19
LONG FORM INFORMATION
1/1/2025
12/31/2025""",
            ),
        ]

        records = extract_standard_schedule_a_records(pages)
        rows = extract_standard_broker_rows(pages)

        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record["carrier_name"], "Standard Life Ins Co of NY")
        self.assertEqual(record["contract_number"], "753370")
        self.assertEqual(record["persons_covered"], "6,159")
        self.assertEqual(record["commission_total"], "0.00")
        self.assertEqual(record["fee_total"], "48,230.78")
        self.assertEqual(
            record["experience_values"]["9a. Premiums: (1) Amount Received"],
            "2,870,399.19",
        )
        self.assertEqual(
            record["experience_values"]["9b(3). Incurred claims (add(1) and (2))"],
            "2,480,723.00",
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].commission_total, "0.00")
        self.assertEqual(rows[0].fee_total, "48,230.78")
        self.assertEqual(rows[0].fee_rows[0].purpose, "Contingent Compensation")

    def test_schedule_a_parser_maps_standard_fields_and_overrides_selected_ftw_schedule(self):
        pages = self._standard_long_form_pages()
        fields = extract_standard_schedule_a_fields(pages)
        mapped = map_extraction_to_rules(
            "test-filing",
            fields,
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )["fields"]
        mapped_by_label = {field.mapped_label: field.proposed_value for field in mapped}

        self.assertEqual(mapped_by_label["1d. Contract/Policy Number"], "168262")
        self.assertEqual(mapped_by_label["1e. Persons Covered (End of Policy Year)"], "63")
        self.assertEqual(mapped_by_label["3b. Amount of Commissions"], "1,506.01")
        self.assertEqual(mapped_by_label["3c. Amount of Fees"], "198.74")
        self.assertEqual(mapped_by_label["9a. Premiums: (1) Amount Received"], "30,312.84")

        summaries = extract_standard_schedule_a_summaries(pages)
        service = FTWilliamsReviewService()
        life_fields = service._fields_with_schedule_a_summary_override(mapped, summaries, "Schedule A-LIFE")
        life_by_label = {field.mapped_label: field.proposed_value for field in life_fields}
        self.assertEqual(life_by_label["1e. Persons Covered (End of Policy Year)"], "107")
        self.assertEqual(life_by_label["3b. Amount of Commissions"], "1,731.10")
        self.assertEqual(life_by_label["3c. Amount of Fees"], "236.46")
        self.assertEqual(life_by_label["9a. Premiums: (1) Amount Received"], "23,945.79")
        self.assertEqual(life_by_label["9a(3). Increase (decrease) in unearned premium reserve"], "-2,104.00")

        ltd_fields = service._fields_with_schedule_a_summary_override(mapped, summaries, "Schedule A-LTD")
        ltd_by_label = {field.mapped_label: field.proposed_value for field in ltd_fields}
        self.assertEqual(ltd_by_label["1e. Persons Covered (End of Policy Year)"], "107")
        self.assertEqual(ltd_by_label["3b. Amount of Commissions"], "1,308.98")
        self.assertEqual(ltd_by_label["3c. Amount of Fees"], "147.90")
        self.assertEqual(ltd_by_label["9a. Premiums: (1) Amount Received"], "12,383.55")
        self.assertEqual(ltd_by_label["9b(2). Increase (decrease) in claim reserves"], "1,610.73")

    def _standard_long_form_pages(self):
        def pdf_text(value):
            return "\n".join(line.strip() for line in value.splitlines())

        def part_i(page, coverage, persons, commission, base, contingent):
            return (
                page,
                pdf_text(f"""
                PAGE: {page}
                (C) PLAN SPONSOR:
                PART I
                1) COVERAGE -
                (a) CARRIER:
                (b) EIN:
                (c) NAIC CODE:
                (d) CONTRACT NUMBER:
                (e) NUMBER OF PERSONS COVERED:
                (f) FROM:
                (g) TO:
                2) INSURANCE FEES AND COMMISSIONS PAID TO AGENTS, BROKERS AND OTHER PERSONS:
                AMOUNT OF COMMISSIONS PAID:
                FEES PAID / AMOUNT:
                TO
                (B) AMOUNT OF COMMISSION PAID FEES PAIDA) NAME & ADDRESS OF AGENT OR
                BROKER TO WHOM COMMISSION OR
                FEES WERE PAID
                COMM. CONT. COMP* GA OVR. (C) AMOUNT (D) PURPOSE
                (E) ORG.
                CODE
                LEAHY CONSULTING SERVICES
                14031 STEEPLESTONE DR
                STE A
                MIDLOTHIAN, VA 23113
                ${base} ${contingent} $0.00 $0.00 3
                TOTAL COMMISSIONS PAID ${base}
                TOTAL CONTINGENT COMP PAID ${contingent}
                TOTAL GA OVERRIDES PAID $0.00
                Standard Insurance Company
                PELLA WINDOW AND DOORS
                12/1/2024
                11/30/2025
                {persons}
                93-0242990
                000-69019
                ${commission}
                $0.00
                THE FINANCIAL DATA BELOW IS PROVIDED FOR YOUR INFORMATION
                IT CAN BE USED TO COMPLETE THE SCHEDULE A FOR THE FORM 5500
                IF YOUR PLAN IS REQUIRED TO FILE SUCH A SCHEDULE
                {coverage}
                PLAN INFORMATION REPORT FOR THE PERIOD OF
                168262
                LONG FORM INFORMATION
                12/1/2024
                11/30/2025
                """),
            )

        def part_iii(page, coverage, premium, values):
            return (
                page,
                pdf_text(f"""
                TO
                PART III -
                7) BENEFIT TYPE:
                EXPERIENCE RATED CONTRACTS
                (a) PREMIUMS: (1) AMOUNT RECEIVED:
                (2) INCREASE (DECREASE) IN DUE BUT UNPAID:
                (3) INCREASE (DECREASE) IN UNEARNED PREMIUM RESERVE:
                (4) EARNED PREMIUM ((1)+(2) - (3)):
                (b) BENEFIT CHARGES: (1) CLAIMS PAID:
                (2) INCREASE (DECREASE) CLAIM RESERVES:
                (3) INCURRED CLAIMS ((1)+(2)):
                (4) CLAIMS CHARGED:
                (c) REMAINDER OF PREMIUM: (1) RETENTION CHARGES:
                PLAN INFORMATION REPORT FOR THE PERIOD OF
                Standard Insurance Company HEREBY CERTIFIES THAT THIS INFORMATION IS COMPLETE AND ACCURATE
                {chr(10).join(values)}
                168262
                {coverage}
                ${premium}
                LONG FORM INFORMATION
                12/1/2024
                11/30/2025
                """),
            )

        return [
            part_i(1, "DENTAL", "63", "1,704.75", "1,506.01", "198.74"),
            part_iii(
                2,
                "DENTAL",
                "30,312.84",
                [
                    "$0.00",
                    "$0.00",
                    "$30,312.84",
                    "$22,882.90",
                    "($78.00)",
                    "$22,804.90",
                    "$22,804.90",
                    "$1,704.75",
                    "$0.00",
                    "$0.00",
                    "$6,601.15",
                    "$682.07",
                    "$639.26",
                    "$0.00",
                    "$9,627.22",
                    "$0.00",
                    "$0.00",
                    "$0.00",
                    "$0.00",
                ],
            ),
            part_i(3, "LIFE INSURANCE", "107", "1,967.56", "1,731.10", "236.46"),
            part_iii(
                4,
                "LIFE INSURANCE",
                "23,945.79",
                [
                    "$112.00",
                    "($2,104.00)",
                    "$26,161.79",
                    "$0.00",
                    "$280.00",
                    "$280.00",
                    "$280.00",
                    "$1,967.56",
                    "$0.00",
                    "$0.00",
                    "$3,388.94",
                    "$588.78",
                    "$2,024.00",
                    "$17,918.72",
                    "$25,888.00",
                    "$0.00",
                    "$0.00",
                    "$0.00",
                    "$0.00",
                ],
            ),
            part_i(5, "LONG TERM DISABILITY", "107", "1,456.88", "1,308.98", "147.90"),
            part_iii(
                6,
                "LONG TERM DISABILITY",
                "12,383.55",
                [
                    "$62.00",
                    "($1,101.00)",
                    "$13,546.55",
                    "$1,200.00",
                    "$1,610.73",
                    "$2,810.73",
                    "$2,810.73",
                    "$1,456.88",
                    "$0.00",
                    "$0.00",
                    "$2,413.15",
                    "$304.83",
                    "$1,495.00",
                    "$5,067.41",
                    "$10,737.27",
                    "$0.00",
                    "$0.00",
                    "$0.00",
                    "$0.00",
                ],
            ),
        ]

    def test_schedule_a_parser_extracts_united_omaha_support_worksheet_pages(self):
        pages = self._united_omaha_pages()

        records = extract_united_omaha_schedule_a_records(pages)
        summaries = extract_united_omaha_schedule_a_summaries(pages)
        rows = extract_united_omaha_broker_rows(pages)

        self.assertEqual(parse_schedule_a_text(pages[0][1]), [])
        by_legacy = {record["legacy_group_id"]: record for record in records}
        self.assertEqual(set(by_legacy), {"GLTD0B432", "GLUG0B432", "GUDH0B432", "GUG0B432"})

        ltd = by_legacy["GLTD0B432"]
        self.assertEqual(ltd["carrier_name"], "United of Omaha Life Insurance Company")
        self.assertEqual(ltd["ein"], "47-0322111")
        self.assertEqual(ltd["naic_code"], "69868")
        self.assertEqual(ltd["group_id"], "G000B432")
        self.assertEqual(ltd["coverage"], "Long Term Disability Insured")
        self.assertEqual(ltd["persons_covered"], "126")
        self.assertEqual(ltd["premium"], "27,628")
        self.assertEqual(ltd["period_begin"], "12/01/2024")
        self.assertEqual(ltd["period_end"], "12/01/2025")

        self.assertEqual(len(summaries), 4)
        self.assertEqual(summaries[0].account_number, "GLTD0B432")
        self.assertEqual(summaries[0].benefit_rows[0].premium, "27,628")
        self.assertEqual(len(rows), 8)
        self.assertEqual(rows[0].name, "GALLAGHER BENEFIT SERVICES INC")
        self.assertEqual(rows[0].commission_total, "4,144")
        self.assertEqual(rows[1].name, "GALLAGHER BENEFIT SERVICES INC")
        self.assertEqual(rows[1].fee_total, "1,424")

    def test_schedule_a_parser_maps_united_omaha_and_overrides_selected_ftw_schedule(self):
        pages = self._united_omaha_pages()
        fields = extract_united_omaha_schedule_a_fields(pages)
        mapped = map_extraction_to_rules(
            "test-filing",
            fields,
            form_type=FormType.SCHEDULE_A,
            source_document_type=DocumentType.SCHEDULE_A,
        )["fields"]
        mapped_by_label = {field.mapped_label: field.proposed_value for field in mapped}

        self.assertEqual(mapped_by_label["1a. Name of Insurance Company"], "United of Omaha Life Insurance Company")
        self.assertEqual(mapped_by_label["1b. Insurance Carrier EIN"], "47-0322111")
        self.assertEqual(mapped_by_label["1c. NAIC Code"], "69868")
        self.assertEqual(mapped_by_label["1d. Contract/Policy Number"], "G000B432")
        self.assertEqual(mapped_by_label["1e. Persons Covered (End of Policy Year)"], "126")
        self.assertEqual(mapped_by_label["1f. Policy Year Beginning Date"], "12/01/2024")
        self.assertEqual(mapped_by_label["1g. Policy Year Ending Date"], "11/30/2025")
        self.assertEqual(mapped_by_label["3b. Amount of Commissions"], "12,851")
        self.assertEqual(mapped_by_label["3c. Amount of Fees"], "4,352")
        self.assertEqual(mapped_by_label["10a. Total premiums or subscription charges paid to carrier"], "85,682")

        combined_rows = extract_united_omaha_combined_broker_rows(pages)
        self.assertEqual(len(combined_rows), 2)
        self.assertEqual(combined_rows[0].commission_total, "12,851")
        self.assertEqual(combined_rows[0].fee_total, "0")
        self.assertEqual(combined_rows[1].commission_total, "0")
        self.assertEqual(combined_rows[1].fee_total, "4,352")
        self.assertEqual(combined_rows[1].name, "GALLAGHER BENEFIT SERVICES INC")

        summaries = extract_united_omaha_schedule_a_summaries(pages)
        service = FTWilliamsReviewService()
        life_fields = service._fields_with_schedule_a_summary_override(mapped, summaries, "Schedule A-LIFE")
        life_by_label = {field.mapped_label: field.proposed_value for field in life_fields}
        self.assertEqual(life_by_label["1d. Contract/Policy Number"], "GLUG0B432")
        self.assertEqual(life_by_label["1e. Persons Covered (End of Policy Year)"], "125")
        self.assertEqual(life_by_label["3b. Amount of Commissions"], "997")
        self.assertEqual(life_by_label["3c. Amount of Fees"], "338")
        self.assertEqual(life_by_label["10a. Total premiums or subscription charges paid to carrier"], "6,649")

        std_fields = service._fields_with_schedule_a_summary_override(mapped, summaries, "Schedule A-STD")
        std_by_label = {field.mapped_label: field.proposed_value for field in std_fields}
        self.assertEqual(std_by_label["1d. Contract/Policy Number"], "GUG0B432")
        self.assertEqual(std_by_label["3b. Amount of Commissions"], "7,008")
        self.assertEqual(std_by_label["3c. Amount of Fees"], "2,375")
        self.assertEqual(std_by_label["10a. Total premiums or subscription charges paid to carrier"], "46,722")

    def _united_omaha_pages(self):
        def pdf_text(value):
            return "\n".join(line.strip() for line in value.splitlines())

        def page(page_number, legacy_group_id, coverage, persons, commission, fee, premium):
            return (
                page_number,
                pdf_text(f"""
                SUPPORT FOR FORM 5500, SCHEDULE A, INSURANCE INFORMATION
                INFORMATION FOR COMPLETION OF PART I
                CAMINO HEALTH CENTER
                SAN JUAN CAPISTRANO, CA
                Name of Carrier: United of Omaha Life Insurance Company - NAIC Code 69868
                EIN Number: 47-0322111
                Group Identification
                Number:
                G000B432 Data for Period: 12-01-2024 to 12-01-2025
                Legacy Group ID: {legacy_group_id}
                Type of Contract: NON-RETENTION
                Benefits Provided Persons Covered
                {coverage} {persons}
                Name of Each Recipient
                Amount of
                Commission
                Paid
                Amount of Service
                Fees Paid or Other
                Fees
                Purpose for
                Which Paid
                Organization
                Type
                GALLAGHER BENEFIT SERVICES INC {commission} Agent or Broker of Record 3
                505 N BRAND BLVD FL 6
                GLENDALE, CA 91203
                GALLAGHER BENEFIT SERVICES INC 0 Other Compensation 3
                NATIONAL INCENTIVE {fee}
                736 S STONE AVE
                LA GRANGE, IL 60525
                INFORMATION FOR COMPLETION OF PART III
                10. Non-experience Rated Contracts:
                Premiums . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . {premium}
                Memo Items: Benefit Charges - Claims Paid . . . . . . . . . . . . . . . . 0
                Administrative Service Fees . . . . . . . . . . . . . . . . . . 0
                Group Office: SOUTHERN CALIFORNIA
                """),
            )

        return [
            page(3, "GLTD0B432", "Long Term Disability Insured", "126", "4,144", "1,424", "27,628"),
            page(4, "GLUG0B432", "Life & AD&D", "125", "997", "338", "6,649"),
            page(5, "GUDH0B432", "Accident only Voluntary", "25", "702", "215", "4,683"),
            page(6, "GUG 0B432", "Short Term Disability Insured", "126", "7,008", "2,375", "46,722"),
        ]

    def test_schedule_a_parser_groups_prudential_same_contract_benefit_pages(self):
        pages = [
            (
                2,
                """
                Insurance Information For SCHEDULE A (Form 5500) Insured Welfare Plan Data
                R.H. White Companies Inc.
                (Item numbers shown correspond to those on Schedule (A) 1 (a) Prudential Insurance Company of America
                1 (b) Prudential's EIN: 22-1211670 1 (c) NAIC code: 68241 1 (d) Contract number or identification: 71492
                1(e) Approximate number 2 Insurance fees and
                of persons covered at end Policy Contract Year commissions paid to
                7 Type of benefit of policy or contract year 1 (f) From 1 (g) To agents or brokers
                1/1/2025 12/31/2025 See Form 27722
                9 Non experience rated contracts:
                a. Total premiums or subscription charges paid to carrier $ 4,759
                GRP 32064 - Rev 1999
                Basic AD&D Insurance 507
                """,
            ),
            (
                3,
                """
                Insurance Information For SCHEDULE A (Form 5500) Insured Welfare Plan Data
                R.H. White Companies Inc.
                (Item numbers shown correspond to those on Schedule (A) 1 (a) Prudential Insurance Company of America
                1 (b) Prudential's EIN: 22-1211670 1 (c) NAIC code: 68241 1 (d) Contract number or identification: 71492
                1/1/2025 12/31/2025 See Form 27722
                9 Non experience rated contracts:
                a. Total premiums or subscription charges paid to carrier $ 37,278
                GRP 32064 - Rev 1999
                Basic Life Insurance 507
                """,
            ),
            (
                4,
                """
                Insurance Information For SCHEDULE A (Form 5500) Insured Welfare Plan Data
                R.H. White Companies Inc.
                (Item numbers shown correspond to those on Schedule (A) 1 (a) Prudential Insurance Company of America
                1 (b) Prudential's EIN: 22-1211670 1 (c) NAIC code: 68241 1 (d) Contract number or identification: 71492
                1/1/2025 12/31/2025 See Form 27722
                9 Non experience rated contracts:
                a. Total premiums or subscription charges paid to carrier $ 69,104
                GRP 32064 - Rev 1999
                Long-Term Disability 198
                """,
            ),
        ]

        fields = extract_prudential_schedule_a_fields(pages)
        by_name = {field.field_name: field.value for field in fields}
        summaries = extract_prudential_schedule_a_summaries(pages)

        self.assertEqual(by_name["1a. Name of Insurance Company"], "Prudential Insurance Company of America")
        self.assertEqual(by_name["1b. Insurance Carrier EIN"], "22-1211670")
        self.assertEqual(by_name["1c. NAIC Code"], "68241")
        self.assertEqual(by_name["1d. Contract/Policy Number"], "71492")
        self.assertEqual(by_name["1e. Persons Covered (End of Policy Year)"], "507")
        self.assertEqual(by_name["10a. Total premiums or subscription charges paid to carrier"], "111,141")
        self.assertEqual(len(summaries), 1)
        self.assertEqual(len(summaries[0].benefit_rows), 3)
        self.assertEqual(summaries[0].benefit_rows[2].benefit_type, "Long-Term Disability")

    def test_schedule_a_parser_extracts_prudential_commission_information_rows(self):
        pages = [
            (
                14,
                """
                ANNUAL REPORT SCHEDULE A(Form 5500) -
                Insurance Information
                (Insured Welfare Plan Commission Information)
                2 Insurance fees and commissions paid to general agents, brokers or other persons:
                71492 RSC INSURANCE BROKERAGE
                INC $45,757
                4TH FLOOR
                160 FEDERAL ST
                BOSTON, MA 2110
                71492 RSC INSURANCE BROKERAGE
                INC
                $12,839
                4TH FLOOR
                160 FEDERAL ST
                BOSTON, MA 2110
                71492 IMG $108
                2960 North Meridian Street
                Indianapolis, IN 46208
                71492 Selman & Company, LLC $13,099
                One Integrity Parkway
                Cleveland, OH 44143
                Includes amounts paid to general agents
                GRP 27722 - Rev 1999 The Prudential Insurance Company of America
                Third Party Administration Fees
                Sales and Service Compensation
                Supplemental Commissions
                Sales and Service Compensation
                """,
            ),
            (15, "12/31/2025\nOrganization \ncode\n3\n3\n5"),
        ]

        rows = extract_prudential_broker_rows(pages)
        by_name = {row.name: row for row in rows}

        self.assertEqual(len(rows), 3)
        self.assertEqual(by_name["RSC INSURANCE BROKERAGE INC"].commission_total, "58,596")
        self.assertEqual(by_name["RSC INSURANCE BROKERAGE INC"].fee_total, "0")
        self.assertEqual(by_name["RSC INSURANCE BROKERAGE INC"].zip_code, "02110")
        self.assertEqual(by_name["IMG"].fee_total, "108")
        self.assertEqual(by_name["IMG"].organization_code, "5")
        self.assertEqual(by_name["Selman & Company, LLC"].fee_total, "13,099")

    def test_prudential_commission_parser_uses_document_contract_instead_of_fixed_sample(self):
        pages = [
            (
                2,
                """
                Insurance Information For SCHEDULE A (Form 5500) Insured Welfare Plan Data
                Homes For The Homeless
                1 (a) Prudential Insurance Company of America
                1 (b) Prudential's EIN: 22-1211670 1 (c) NAIC code: 68241 1 (d) Contract number or identification: 15408
                1/1/2025 12/31/2025 See Form 27722
                9 Non experience rated contracts:
                a. Total premiums or subscription charges paid to carrier $ 2,031
                GRP 32064 - Rev 1999
                Basic AD&D Insurance 356
                """,
            ),
            (
                7,
                """
                ANNUAL REPORT SCHEDULE A(Form 5500) - Insurance Information
                (Insured Welfare Plan Commission Information)
                15408 RSC INSURANCE BROKERAGE INC $5,919
                4TH FLOOR
                160 FEDERAL ST
                BOSTON, MA 2110
                15408 RSC INSURANCE BROKERAGE INC $2,936
                4TH FLOOR
                160 FEDERAL ST
                BOSTON, MA 2110
                15408 IMG $76
                2960 North Meridian Street
                Indianapolis, IN 46208
                15408 Selman & Company, LLC $2,938
                One Integrity Parkway
                Cleveland, OH 44143
                Includes amounts paid to general agents
                GRP 27722 - Rev 1999 The Prudential Insurance Company of America
                """,
            ),
        ]

        rows = extract_prudential_broker_rows(pages)
        by_name = {row.name: row for row in rows}

        self.assertEqual(len(rows), 3)
        self.assertEqual(by_name["RSC INSURANCE BROKERAGE INC"].commission_total, "8,855")
        self.assertEqual(by_name["IMG"].fee_total, "76")
        self.assertEqual(by_name["Selman & Company, LLC"].fee_total, "2,938")

    def test_schedule_a_parser_extracts_summary_table_broker_rows(self):
        pages = [
            (
                1,
                """
                Commissions
                Total commissions
                The following figure represents commissions that are to be reported on Schedule A, Line 3, Element (b):
                Contract ID Contract name Commissions paid
                000FG530 PREFERRED BENEFITS GROUP $16,464.93
                Total commissions for plan $16,464.93

                Group insurance coverages Commissions paid
                AD&D 736.11
                Dental (Insured) 3125.84
                Life 2785.21

                Fees
                Total fees
                The following figure represents fees that are to be reported on Schedule A, Line 3, Element (c):
                Contract ID Contract name Amount
                000F5894 NFP CORPORATE SERVICES NY LLC $2,645.25
                Total Fees Paid $2,645.25

                Group insurance coverages Gross premium paid
                Vision (Insured) $8,796.05
                Total premium paid $186,242.73
                """,
            )
        ]

        rows = extract_summary_table_broker_rows(pages)
        by_name = {row.name: row for row in rows}

        self.assertEqual(set(by_name), {"PREFERRED BENEFITS GROUP", "NFP CORPORATE SERVICES NY LLC"})
        self.assertEqual(by_name["PREFERRED BENEFITS GROUP"].commission_total, "16,464.93")
        self.assertEqual(by_name["PREFERRED BENEFITS GROUP"].organization_code, "3")
        self.assertEqual(by_name["NFP CORPORATE SERVICES NY LLC"].fee_total, "2,645.25")
        self.assertEqual(by_name["NFP CORPORATE SERVICES NY LLC"].organization_code, "3")

    def test_schedule_a_parser_merges_summary_table_broker_sections(self):
        pages = [
            (
                1,
                """
                The following figure represents commissions that are to be reported on Schedule A, Line 3, Element (b):
                Contract ID Contract name Commissions paid
                000FG530 PREFERRED BENEFITS GROUP $10.00

                The following figure represents fees that are to be reported on Schedule A, Line 3, Element (c):
                Contract ID Contract name Amount
                000FG530 PREFERRED BENEFITS GROUP $2.50
                """,
            )
        ]

        rows = extract_summary_table_broker_rows(pages)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].name, "PREFERRED BENEFITS GROUP")
        self.assertEqual(rows[0].commission_total, "10")
        self.assertEqual(rows[0].fee_total, "2.50")


if __name__ == "__main__":
    unittest.main()
